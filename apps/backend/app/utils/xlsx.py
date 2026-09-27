"""Geração de .xlsx dos relatórios exportáveis (cabeçalho, filtros, larguras)
e proteção contra injeção de fórmula nas células."""

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
DATETIME_FORMAT = "DD/MM/YYYY HH:MM:SS"


def formula_safe(value: str) -> str:
    """Neutraliza injeção de fórmula (CWE-1236): nome/valor/telefone vêm da
    planilha importada por qualquer usuário e, sem isso, um valor como
    "=cmd|'/c calc'!A0" seria executado ao abrir o relatório no Excel/
    LibreOffice — o openpyxl trata string começando com "=" como fórmula
    igual ao próprio Excel."""

    if value and value[0] in FORMULA_PREFIXES:
        return "'" + value
    return value


def texto_nunca_formula(cell) -> None:
    """Texto que começa com "=" é gravado como texto, não como fórmula: o
    relatório não executa nada ao abrir e o conteúdo aparece igual (sem o
    apóstrofo que formula_safe acrescenta)."""

    if cell.data_type == "f":
        cell.data_type = "s"


def build_xlsx(headers: list[str], rows: list[list], formatos: dict[int, str] | None = None) -> bytes:
    """Gera um .xlsx com cabeçalho destacado, painel congelado, autofiltro e
    largura de coluna ajustada — usado por todos os relatórios exportáveis.
    `formatos`: índice da coluna (0 = primeira) → formato de número."""

    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="003090")
    for row in rows:
        ws.append(row)
        for cell in ws[ws.max_row]:
            texto_nunca_formula(cell)
            if isinstance(cell.value, datetime):
                cell.number_format = DATETIME_FORMAT
            elif isinstance(cell.value, date):
                cell.number_format = "DD/MM/YYYY"
            if formatos and cell.column - 1 in formatos:
                cell.number_format = formatos[cell.column - 1]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, header in enumerate(headers, start=1):
        widths = [len(header)] + [len(str(row[i - 1])) for row in rows if row[i - 1] is not None]
        ws.column_dimensions[get_column_letter(i)].width = min(max(widths) + 4, 40)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
