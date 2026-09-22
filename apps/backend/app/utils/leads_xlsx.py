"""Exportação de leads cobrados para .xlsx — transforma objetos Lead em bytes
de planilha pronta para importar em outro sistema."""

import io
import re
from typing import Iterable

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from .phone import is_valid_phone, normalize_phone

# Cor de preenchimento do cabeçalho definida pelo negócio
_COR_CABECALHO = "FCE4D6"

_CABECALHO = ("Codigo", "Nome", "CPF", "Celular")
_LARGURAS = (12, 24, 16, 18)


def formatar_codigo(valor: str | None) -> str:
    if not valor:
        return ""
    digitos = re.sub(r"\D", "", str(valor))
    if not digitos:
        return ""
    if len(digitos) <= 8:
        return digitos.zfill(8)
    return digitos


def primeiro_nome(nome: str | None) -> str:
    """Normaliza para o primeiro nome capitalizado; vazio/None/sem letra → ''."""
    if not nome:
        return ""
    # Divide por qualquer caractere que não seja letra (incluindo acentuadas),
    # hífen ou apóstrofo — dígitos, pontuação e _ são separadores
    partes = re.split(r"[^\w\-']|[0-9_]", nome, flags=re.UNICODE)
    # Pega o primeiro pedaço não vazio (pode ser "ANA-LUIZA" ou "D'AVILA" inteiro)
    primeiro = next((p for p in partes if p), "")
    if not primeiro:
        return ""
    # Remove hífens/apóstrofos soltos nas bordas
    primeiro = primeiro.strip("-'")
    # Garante que sobrou ao menos uma letra
    if not re.search(r"[^\W\d_]", primeiro, flags=re.UNICODE):
        return ""
    # Capitaliza cada segmento separado por hífen ou apóstrofo
    resultado = re.sub(
        r"[^-']+",
        lambda m: m.group(0).capitalize(),
        primeiro,
        flags=re.UNICODE,
    )
    return resultado


def formatar_cpf(valor: str | None) -> str:
    if not valor:
        return ""
    digitos = re.sub(r"\D", "", str(valor))
    if not digitos:
        return ""
    n = len(digitos)
    if n <= 11:
        d = digitos.zfill(11)
        return f"{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}"
    elif n <= 14:
        d = digitos.zfill(14)
        return f"{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}"
    # Mais de 14 dígitos: devolve sem máscara
    return digitos


def formatar_celular(valor: str | None) -> str:
    if not valor:
        return ""
    if is_valid_phone(valor):
        return normalize_phone(valor)
    return ""


def linhas_para_exportacao(leads: Iterable) -> list[tuple[str, str, str, str]]:
    resultado = []
    for lead in leads:
        codigo = formatar_codigo(lead.codigo_cliente)
        nome = primeiro_nome(lead.nome)
        cpf = formatar_cpf(lead.cpf)
        celular = formatar_celular(lead.celular)
        resultado.append((codigo, nome, cpf, celular))
    return resultado


def gerar_xlsx_leads(leads: Iterable) -> bytes:
    """Monta a planilha Excel com os leads formatados e devolve os bytes do arquivo."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"

    # Estilo do cabeçalho: preenchimento laranja claro, Calibri 11 sem negrito
    fill = PatternFill(fill_type="solid", fgColor=_COR_CABECALHO)
    font = Font(name="Calibri", size=11, bold=False, color="000000")
    alinhamento = Alignment(horizontal="left")

    for col_idx, titulo in enumerate(_CABECALHO, start=1):
        cell = ws.cell(row=1, column=col_idx, value=titulo)
        cell.fill = fill
        cell.font = font
        cell.alignment = alinhamento

    # Larguras de coluna (openpyxl usa letras A, B, C, D)
    letras = ("A", "B", "C", "D")
    for letra, largura in zip(letras, _LARGURAS):
        ws.column_dimensions[letra].width = largura

    # Painel congelado abaixo do cabeçalho
    ws.freeze_panes = "A2"

    # Contador local evita o max_row quadrático do openpyxl
    row_idx = 1
    for linha_dados in linhas_para_exportacao(leads):
        row_idx += 1
        for col_idx, valor in enumerate(linha_dados, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=str(valor))
            cell.number_format = "@"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
