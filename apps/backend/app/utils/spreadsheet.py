"""Geração do modelo de planilha por faixa e leitura de planilhas enviadas —
tudo em .xlsx (openpyxl), sem depender de pandas."""

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

REQUIRED_COLUMNS = ["codigo_cliente", "nome", "celular", "valor"]


def build_model_columns(variable_internal_names: list[str]) -> list[str]:
    columns = list(REQUIRED_COLUMNS)
    for name in variable_internal_names:
        if name not in columns:
            columns.append(name)
    return columns


def build_model_xlsx(variable_internal_names: list[str]) -> bytes:
    """Gera o .xlsx de modelo já formatado (cabeçalho em destaque e coluna
    com largura legível) para o usuário baixar, preencher e subir de volta."""

    columns = build_model_columns(variable_internal_names)
    wb = Workbook()
    ws = wb.active
    ws.title = "modelo"
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="003090")
    ws.freeze_panes = "A2"
    for i, column in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(len(column) + 4, 14)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _require_xlsx(filename: str) -> None:
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Envie um arquivo .xlsx")


def parse_uploaded_spreadsheet(filename: str, content: bytes) -> list[dict[str, str]]:
    """Retorna uma lista de linhas como dicts {coluna: valor}, a partir de um
    .xlsx enviado pelo usuário."""

    _require_xlsx(filename)
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    result = []
    for row in rows[1:]:
        if row is None or all(cell is None for cell in row):
            continue
        record = {
            headers[i]: ("" if row[i] is None else str(row[i]).strip())
            for i in range(len(headers))
            if i < len(row)
        }
        result.append(record)
    return result


def read_spreadsheet_headers(filename: str, content: bytes) -> list[str]:
    """Lê só a primeira linha (cabeçalho) de um .xlsx enviado, para que o
    usuário escolha, em uma lista suspensa, qual coluna real vira cada
    variável/campo — sem precisar que a planilha use nomes fixos."""

    _require_xlsx(filename)
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    first_row = next(ws.iter_rows(values_only=True), None)
    if not first_row:
        return []
    return [str(h).strip() if h is not None else "" for h in first_row if h is not None and str(h).strip()]
