"""Campanhas de cobrança (tela Campanhas): uma base filtrada com os mesmos
filtros da Cobrança, templates próprios e um período em que roda sozinha.

Cada campanha tem uma faixa só dela ("Campanha: …"), onde ficam número,
template e variáveis, como no remarketing. Com "Envio automático", roda todo
dia de disparo dentro do período (um dia só = início e fim iguais); sem ele,
só pelo "Colocar na fila agora". No dia,
antes do horário de início, o agendador busca a base da campanha no SETA
(só clientes em atraso, ou na faixa só de campanhas "Antecipado", quando a
campanha a escolhe) e coloca na fila quem ainda não recebeu. Com uma
planilha de clientes, a base fica restrita a eles e, se a campanha usa os
valores da planilha, valor, celular e colunas das variáveis vêm dela. Com
"Todos os clientes da planilha", entra todo cliente dela com parcela em aberto
no SETA, em atraso ou não (`faixa_na_campanha`). Daí em diante é o fluxo
normal: blacklist, uma cobrança por cliente por dia (a não ser na campanha
com "Incluir quem já recebeu mensagem hoje"), pausas, expiração no fim da
janela.

Quem entra na fila da campanha vira também um Lead da sua faixa de atraso
marcado com a campanha (é o que a Efetividade e a exportação de leads filtram
por campanha). Num dia de disparo seguinte ao envio da campanha, o cliente
entra também na fila da régua da faixa de atraso em que estiver naquele dia
(`enfileirar_na_regua`): no mesmo dia não dá, só vai uma cobrança por dia.
"""

import io
import logging
import re
from datetime import date, datetime, timedelta
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy.orm import Session, selectinload

from . import cobranca_base, lojas as lojas_base, models, pausas, seta_client
from .elegibilidade import STATUS_OCUPA_CLIENTE, clientes_bloqueados_hoje, ocupa_cliente, travar_entrada_na_fila
from .fila_automatica import (
    enfileirar_clientes,
    enfileirar_leads,
)
from .leads_service import gerar_leads_de_clientes
from .regras_db import carregar_regras
from .timezone import hoje_br, inicio_do_dia_utc, inicio_hoje_utc, para_br
from .utils.phone import is_valid_phone
from .utils.valor import ler_valor
from .variaveis_template import CAMPOS_CLIENTE, extrair_placeholders, formatar_moeda

logger = logging.getLogger("campanhas")

PREFIXO_FAIXA = "Campanha: "
MAX_LINHAS_PLANILHA = 200_000

# Filtros de loja: viram uma lista de códigos de filial antes da consulta.
_FILTROS_LOJA = ("loja", "regional", "estado", "cluster_inad", "cluster_populacao", "cobradora")


def nome_faixa(nome: str) -> str:
    return PREFIXO_FAIXA + nome


def em_periodo(campanha: models.Campanha, hoje: date) -> bool:
    if campanha.modo == "unica":  # legado: período de um dia
        return campanha.data_inicio == hoje
    if campanha.data_inicio and hoje < campanha.data_inicio:
        return False
    if campanha.data_fim and hoje > campanha.data_fim:
        return False
    return True


def filtros_para_busca(
    db: Session, filtros: dict, clientes: list[str] | None, *, todos_da_planilha: bool = False
) -> dict:
    """Argumentos de `cobranca_base.buscar_base` a partir dos filtros salvos
    (mesmos nomes dos parâmetros de /cobranca/clientes). Numa campanha, o
    padrão é pegar todos os dias da faixa e ignorar a matriz do WhatsApp:
    quem decide quem entra são os filtros da própria campanha. Com
    `todos_da_planilha` (e a planilha), vale qualquer dia de atraso."""

    def lista(chave: str) -> list[str] | None:
        return list(filtros.get(chave) or []) or None

    def data(chave: str) -> date | None:
        valor = filtros.get(chave)
        return date.fromisoformat(valor) if valor else None

    def decimal(chave: str) -> Decimal | None:
        valor = filtros.get(chave)
        return Decimal(str(valor)) if valor not in (None, "") else None

    codigos_loja = lojas_base.combinar_lojas(
        db,
        lista("loja"),
        regional=lista("regional"),
        estado=lista("estado"),
        cluster_inad=lista("cluster_inad"),
        cluster_populacao=lista("cluster_populacao"),
        cobradora=lista("cobradora"),
    )
    return dict(
        apenas_primeiro_dia=bool(filtros.get("apenas_primeiro_dia", False)),
        somente_regra_whatsapp=bool(filtros.get("somente_regra_whatsapp", False)),
        faixas=lista("faixa"),
        clusters=lista("cluster"),
        faixas_compra=lista("faixa_compra"),
        lojas=codigos_loja,
        status_cliente=lista("status_cliente"),
        restricoes_spc=lista("restricao_spc"),
        vencimento_de=data("vencimento_de"),
        vencimento_ate=data("vencimento_ate"),
        valor_atraso_min=decimal("valor_atraso_min"),
        valor_atraso_max=decimal("valor_atraso_max"),
        valor_atraso_com_juros=bool(filtros.get("valor_atraso_com_juros", False)),
        codigos=list(clientes) if clientes else None,
        incluir_so_campanhas=True,
        todos_os_dias=bool(todos_da_planilha and clientes),
    )


# Campo de template que não existe para quem não está em atraso: bloqueado na
# campanha da faixa só de campanhas (Antecipado).
CAMPOS_SO_EM_ATRASO = ("valor_atraso",)


def faixa_na_campanha(regras, dias_atraso: int) -> str | None:
    """Faixa de quem entra pela campanha com todos os clientes da planilha:
    em atraso, a faixa de atraso cadastrada ou, se nenhuma cobre o dia, uma
    faixa só com ele (ex.: "1"); sem atraso, a faixa só de campanhas que cobre
    o dia (ou a primeira delas, ex.: "Antecipado"). None se não há faixa só de
    campanhas cadastrada para quem não está em atraso."""

    if dias_atraso >= 1:
        faixa = regras.faixa_por_dias(dias_atraso)
        return faixa if faixa and faixa not in regras.nomes_faixa_so_campanhas else str(dias_atraso)
    faixa = regras.faixa_por_dias(dias_atraso)
    if faixa in regras.nomes_faixa_so_campanhas:
        return faixa
    return next(iter(regras.nomes_faixa_so_campanhas), None)


def faixa_so_campanhas(db: Session, faixas: list[str]) -> str | None:
    """A faixa só de campanhas (ex.: "Antecipado") entre as faixas escolhidas."""
    so_campanhas = carregar_regras(db).nomes_faixa_so_campanhas
    return next((f for f in faixas if f in so_campanhas), None)


def erro_faixa_so_campanhas(db: Session, faixas: list[str], *, valor_atraso: bool, mapeamentos) -> str | None:
    """Mensagem de erro se a campanha usa a faixa só de campanhas com outra
    faixa, com filtro de valor em atraso ou com variável "Valor em atraso"."""

    faixa = faixa_so_campanhas(db, faixas)
    if faixa is None:
        return None
    if len(faixas) > 1:
        return f"A faixa {faixa} não pode ser combinada com outras faixas"
    if valor_atraso:
        return f"A faixa {faixa} não tem valor em atraso; tire esse filtro"
    if bloqueado := _campo_so_em_atraso(mapeamentos):
        return f'A faixa {faixa} não tem "{CAMPOS_CLIENTE[bloqueado]}"; troque essa variável do template'
    return None


def _campo_so_em_atraso(mapeamentos) -> str | None:
    for m in mapeamentos:
        campos = [m.column_name] if m.fonte_tipo == "campo_cliente" else (
            extrair_placeholders(m.expressao or "") if m.fonte_tipo == "expressao" else []
        )
        if bloqueado := next((c for c in campos if c in CAMPOS_SO_EM_ATRASO), None):
            return bloqueado
    return None


def erro_todos_da_planilha(db: Session, faixas: list[str], *, valor_atraso: bool, mapeamentos) -> str | None:
    """Mensagem de erro se a campanha com todos os clientes da planilha tem
    filtro de faixa ou de valor em atraso, usa a variável "Valor em atraso"
    (quem não está em atraso não tem) ou não há faixa só de campanhas."""

    if faixas:
        return "Com todos os clientes da planilha, a campanha não filtra por faixa; tire esse filtro"
    if valor_atraso:
        return "Com todos os clientes da planilha, tire o filtro de valor em atraso (quem não está em atraso não tem)"
    if bloqueado := _campo_so_em_atraso(mapeamentos):
        return (
            f'Com todos os clientes da planilha, quem não está em atraso não tem "{CAMPOS_CLIENTE[bloqueado]}"; '
            "troque essa variável do template"
        )
    if not carregar_regras(db).nomes_faixa_so_campanhas:
        return "Cadastre uma faixa só de campanhas (ex.: Antecipado) para quem não está em atraso"
    return None


def ja_receberam(db: Session, campanha: models.Campanha, agora: datetime | None = None) -> set[str]:
    """Quem já está na fila ou recebeu desta campanha (em qualquer data, ou
    dentro do prazo de recontato, quando a recorrente tem um). Quem entrou e
    não foi cobrado (erro, parado, descartado) pode entrar de novo."""

    q = db.query(models.QueueItem.codigo_cliente).filter(
        models.QueueItem.faixa_id == campanha.faixa_id,
        ocupa_cliente(),
    )
    if campanha.recontato_dias:
        # Em dias de calendário (GMT-3), não em janelas de 24 h: com recontato
        # de 7 dias, quem recebeu no dia 1 volta no dia 8, seja qual for a hora
        dia = para_br(agora).date() if agora else hoje_br()
        q = q.filter(
            models.QueueItem.created_at >= inicio_do_dia_utc(dia - timedelta(days=campanha.recontato_dias - 1))
        )
    return {codigo for (codigo,) in q}


def selecionar(db: Session, campanha: models.Campanha) -> dict:
    """{"status": "ready", "clientes": [...], "total_base": n} com quem entraria
    hoje (base filtrada menos quem já recebeu, blacklist e sem decisão), ou
    {"status": "processing"} enquanto o SETA calcula a base em segundo plano."""

    resultado = diagnosticar(db, campanha)
    if resultado["status"] != "ready":
        return resultado

    return {
        "status": "ready",
        "total_base": resultado["total_base"],
        "clientes": resultado["clientes"]
    }


_COLUNAS_VALOR = ("VALOR", "VALOR EM ATRASO", "VALOR A COBRAR", "VALOR ATRASO", "VALOR EM ABERTO", "VALOR DIVIDA")
_COLUNAS_CELULAR = ("CELULAR", "TELEFONE", "FONE", "WHATSAPP")


def _coluna_da_linha(linha: dict[str, str], nomes: tuple[str, ...]) -> str | None:
    for coluna, valor in linha.items():
        if lojas_base._normalizar(coluna) in nomes and str(valor).strip():
            return str(valor).strip()
    return None


def _decimal_planilha(texto: str) -> Decimal | None:
    """Valor da planilha maior que zero; zero, negativo ou sem número fica
    com o valor do SETA."""
    valor = ler_valor(texto)
    return valor if valor is not None and valor > 0 else None


def _colunas_para_variaveis(linha: dict[str, str]) -> dict[str, str]:
    """Linha da planilha como contexto das variáveis: coluna de valor sai
    formatada como a mensagem mostra ("1.234,50", sem "R$")."""

    contexto = dict(linha)
    for coluna, texto in linha.items():
        if lojas_base._normalizar(coluna) in _COLUNAS_VALOR and (valor := ler_valor(str(texto))) is not None:
            contexto[coluna] = formatar_moeda(valor)
    return contexto


def conferir_obrigatorios(linhas_por_codigo: dict[str, dict], cadastro: dict[str, dict]) -> tuple[dict[str, dict], dict[str, int], dict[str, int]]:
    """Aplica as regras R1 a R3 aos clientes importados."""
    from .utils.document import format_cpf

    _COLUNAS_NOME = ("NOME", "NOME CLIENTE", "CLIENTE")
    _COLUNAS_CPF = ("CPF", "CPF/CNPJ", "CPFCNPJ", "CPF CNPJ", "DOCUMENTO")
    _COLUNAS_CELULAR = ("CELULAR", "TELEFONE", "FONE", "WHATSAPP")
    ORDEM_TELEFONES = ("telefone2", "telefone4", "telefone3", "telefone1")

    def _coluna(linha: dict, nomes: tuple) -> str:
        for c, v in linha.items():
            if lojas_base._normalizar(c) in nomes and str(v).strip():
                return str(v).strip()
        return ""

    planilha_linhas = {}
    pendentes: dict[str, int] = {}
    completados_pelo_seta: dict[str, int] = {"nome": 0, "cpf": 0, "celular": 0}
    
    for codigo, linha in linhas_por_codigo.items():
        cad = cadastro.get(codigo)
        if not cad:
            pendentes["código não encontrado no SETA"] = pendentes.get("código não encontrado no SETA", 0) + 1
            continue

        origem = {}
        
        # NOME
        nome_seta = (cad.get("nome") or "").strip()
        nome_planilha = _coluna(linha, _COLUNAS_NOME)
        if nome_seta:
            nome_final = nome_seta
            origem["nome"] = "SETA"
            if not nome_planilha:
                completados_pelo_seta["nome"] += 1
        elif nome_planilha:
            nome_final = nome_planilha
            origem["nome"] = "Planilha"
        else:
            pendentes["sem nome"] = pendentes.get("sem nome", 0) + 1
            continue
            
        # CPF
        cpf_seta = (cad.get("cpfcnpj") or "").strip()
        cpf_planilha = _coluna(linha, _COLUNAS_CPF)
        if cpf_seta:
            cpf_final = cpf_seta
            origem["cpf"] = "SETA"
            if not cpf_planilha:
                completados_pelo_seta["cpf"] += 1
        elif cpf_planilha:
            cpf_final = format_cpf(cpf_planilha)
            origem["cpf"] = "Planilha"
        else:
            pendentes["sem CPF"] = pendentes.get("sem CPF", 0) + 1
            continue
            
        # CELULAR
        celular_planilha = _coluna(linha, _COLUNAS_CELULAR)
        celular_final = ""
        if celular_planilha and is_valid_phone(celular_planilha):
            celular_final = celular_planilha
            origem["celular"] = "Planilha"
        else:
            for tel in ORDEM_TELEFONES:
                c = (cad.get(tel) or "").strip()
                if c and is_valid_phone(c):
                    celular_final = c
                    origem["celular"] = "SETA"
                    completados_pelo_seta["celular"] += 1
                    break
                    
        if not celular_final:
            pendentes["sem telefone válido"] = pendentes.get("sem telefone válido", 0) + 1
            continue
            
        nova_linha = dict(linha)
        nova_linha["_nome"] = nome_final
        nova_linha["_cpf"] = cpf_final
        nova_linha["_celular"] = celular_final
        nova_linha["_origem"] = origem
        planilha_linhas[codigo] = nova_linha

    completados_pelo_seta = {k: v for k, v in completados_pelo_seta.items() if v > 0}
    return planilha_linhas, pendentes, completados_pelo_seta


def colunas_em_branco(campanha: models.Campanha, mapeamentos: list[models.FaixaVariableMapping]) -> list[dict]:
    linhas = list((campanha.planilha_linhas or {}).values())
    if not linhas:
        return []
        
    resultado = []
    for m in mapeamentos:
        if m.fonte_tipo != "coluna" or not m.column_name:
            continue
            
        col_name = m.column_name
        qtd_afetados = 0
        for linha in linhas:
            if not str(linha.get(col_name, "")).strip():
                qtd_afetados += 1
                
        if qtd_afetados > 0:
            sugestao = None
            norm_col = lojas_base._normalizar(col_name)
            for k, v in CAMPOS_CLIENTE.items():
                if lojas_base._normalizar(k) == norm_col or lojas_base._normalizar(v) == norm_col:
                    sugestao = k
                    break
            
            resultado.append({
                "variavel_id": m.template_variable_id,
                "coluna_nome": col_name,
                "qtd_afetados": qtd_afetados,
                "sugestao_campo": sugestao
            })
            
    return resultado


def valor_da_planilha(linha: dict) -> Decimal | None:
    """Valor da linha da planilha; None quando não é válido (fica o do SETA,
    se o usuário autorizou em `Campanha.valor_seta_autorizado`)."""

    return _decimal_planilha(_coluna_da_linha(linha, _COLUNAS_VALOR) or "")


def valor_invalido(campanha: models.Campanha) -> int:
    """Clientes da planilha sem valor válido."""

    return sum(1 for linha in (campanha.planilha_linhas or {}).values() if valor_da_planilha(linha) is None)


def com_dados_da_planilha(cliente: dict, campanha: models.Campanha) -> dict:
    """Cliente do SETA com valor, nome, CPF e celular conferidos pela planilha."""
    linha = (campanha.planilha_linhas or {}).get(cliente["codigo"])
    if not linha:
        return cliente
    novo = dict(cliente)
    if (valor := valor_da_planilha(linha)) is not None:
        novo["valor_cobrar"] = valor
        novo["valor_atraso"] = valor
    if "_nome" in linha:
        novo["nome"] = linha["_nome"]
    if "_cpf" in linha:
        novo["cpfcnpj"] = linha["_cpf"]
    if "_celular" in linha:
        novo["celular"] = linha["_celular"]
        novo["celular_original"] = linha["_celular"]
    return novo


def _carregar_faixa(db: Session, campanha: models.Campanha) -> models.Faixa:
    return (
        db.query(models.Faixa)
        .options(
            selectinload(models.Faixa.envios).selectinload(models.FaixaEnvio.template).selectinload(
                models.Template.variables
            ),
            selectinload(models.Faixa.variable_mappings),
        )
        .filter(models.Faixa.id == campanha.faixa_id)
        .one()
    )


def _bloqueados(db: Session, campanha: models.Campanha) -> set[str]:
    """Quem não pode entrar na fila desta campanha agora. A campanha fora da
    regra de uma cobrança por dia ignora pendentes e cobrados de hoje em
    outros caminhos (não duplicar na própria campanha é com `ja_receberam`)."""

    if campanha.incluir_cobrados_hoje:
        travar_entrada_na_fila(db)
        return set()
    return clientes_bloqueados_hoje(db)


def executar(db: Session, campanha: models.Campanha, *, created_by: str | None = None) -> dict:
    """Busca a base e coloca na fila. Devolve {"status": "processing"} se a
    base ainda está sendo calculada (quem chamou tenta de novo depois) ou o
    resumo {"encontrados", "na_fila"}, que fica gravado na campanha."""

    selecao = selecionar(db, campanha)
    if selecao["status"] != "ready":
        return selecao
    clientes = selecao["clientes"]
    extras = None
    if campanha.fonte_valores == "planilha":
        clientes = [com_dados_da_planilha(c, campanha) for c in clientes]
        extras = {codigo: _colunas_para_variaveis(linha) for codigo, linha in (campanha.planilha_linhas or {}).items()}
    na_fila = 0
    if clientes:
        juros = carregar_regras(db).juros
        parcelas = seta_client.buscar_parcelas_cobranca([c["codigo"] for c in clientes], juros=juros)
        enfileirados: list[dict] = []
        na_fila = enfileirar_clientes(
            db,
            _carregar_faixa(db, campanha),
            clientes,
            bloqueados=_bloqueados(db, campanha),
            juros=juros,
            parcelas=parcelas,
            origem="na campanha",
            colunas_extras=extras,
            enfileirados=enfileirados,
        )
        db.flush()
        # Lead na faixa de atraso do cliente, marcado com a campanha: o envio
        # o marca como cobrado (Efetividade/exportação) e leva o cliente à régua.
        gerar_leads_de_clientes(
            db, enfileirados, created_by=created_by, campanha_id=campanha.id, parcelas=parcelas
        )
    resumo = {"status": "ready", "encontrados": len(clientes), "na_fila": na_fila}
    campanha.ultima_execucao = datetime.utcnow()
    campanha.parada_em = None
    campanha.ultimo_resultado = {"encontrados": len(clientes), "na_fila": na_fila}
    db.commit()
    logger.info("Campanha '%s': %s", campanha.nome, resumo)
    return resumo


def enfileirar_na_regua(db: Session) -> dict:
    """Quem recebeu mensagem de campanha em dia anterior entra na fila da
    régua da faixa de atraso em que está hoje (base do SETA recalculada: quem
    pagou ou saiu do atraso não entra). Cliente bloqueado hoje (já na fila ou
    cobrado em outra faixa) tenta de novo no próximo dia de disparo.
    Devolve {"status": "processing"} enquanto a base calcula, ou o resumo."""

    leads = (
        db.query(models.Lead)
        .filter(
            models.Lead.campanha_id != "",
            models.Lead.status == "cobrado",
            models.Lead.regua_em.is_(None),
            models.Lead.cobrado_em < inicio_hoje_utc(),
        )
        .all()
    )
    if not leads:
        return {"status": "ready", "clientes": 0, "na_fila": 0}
    codigos = sorted({l.codigo_cliente for l in leads})
    job = cobranca_base.buscar_base(db, apenas_primeiro_dia=False, somente_regra_whatsapp=False, codigos=codigos)
    if job["status"] != "ready":
        return {"status": "processing"}
    base = [c for c in job["data"] if c["dias_atraso"] >= 1]
    na_fila = 0
    if base:
        gerar_leads_de_clientes(db, base, created_by=None)
        na_fila = enfileirar_leads(db, base)

    faixa_por_id = {f.id: f.name for f in db.query(models.Faixa.id, models.Faixa.name)}
    na_regua_hoje = {
        (codigo, faixa_por_id.get(faixa_id))
        for codigo, faixa_id in db.query(models.QueueItem.codigo_cliente, models.QueueItem.faixa_id).filter(
            models.QueueItem.codigo_cliente.in_(codigos),
            models.QueueItem.created_at >= inicio_hoje_utc(),
            # Erro também resolve o dia: sem isso, o item com variável sem valor
            # era recriado a cada ciclo do worker (5 s) até o fim da janela.
            models.QueueItem.status.in_((*STATUS_OCUPA_CLIENTE, models.QueueStatus.error)),
        )
    }
    com_envio = {
        f.name
        for f in db.query(models.Faixa).filter(models.Faixa.active.is_(True), models.Faixa.tipo == models.TIPO_REGUA)
        if any(e.active and e.template_id for e in f.envios)
    }
    faixa_hoje = {c["codigo"]: c["faixa"] for c in base}
    agora = datetime.utcnow()
    for lead in leads:
        faixa = faixa_hoje.get(lead.codigo_cliente)
        # Resolvido: saiu da base (pagou, blacklist…), a faixa não tem régua
        # ligada ou o cliente já está na fila da régua hoje.
        if faixa is None or faixa not in com_envio or (lead.codigo_cliente, faixa) in na_regua_hoje:
            lead.regua_em = agora
    db.commit()
    resumo = {"status": "ready", "clientes": len(codigos), "na_fila": na_fila}
    logger.info("Régua de quem recebeu campanha: %s", resumo)
    return resumo


def campanhas_para_hoje(db: Session) -> list[models.Campanha]:
    """Com envio automático, não arquivadas nem pausadas, dentro do período,
    ainda não rodadas hoje e com algum número + template ativo (ligar antes de
    atribuir não perde o dia)."""

    hoje = hoje_br()
    pausadas = {p.valor for p in pausas.ativas(db) if p.escopo == "faixa"}
    return [
        c
        for c in db.query(models.Campanha).filter(
            models.Campanha.ativa.is_(True), models.Campanha.arquivada_em.is_(None)
        )
        if em_periodo(c, hoje)
        and c.ultima_execucao_dia != hoje
        and c.faixa_id not in pausadas
        and any(e.active for e in c.faixa.envios)
    ]


# --- Leitura de planilhas -----------------------------------------------------


def _linhas_xlsx(filename: str, content: bytes) -> list[list[str]]:
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Envie um arquivo .xlsx")
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - arquivo corrompido ou não é xlsx de verdade
        raise ValueError("Não consegui abrir a planilha; confira se é um .xlsx válido") from exc
    linhas = []
    for linha in wb.active.iter_rows(values_only=True):
        if len(linhas) > MAX_LINHAS_PLANILHA:
            raise ValueError(f"A planilha passa de {MAX_LINHAS_PLANILHA} linhas")
        linhas.append(["" if c is None else _texto_celula(c) for c in linha])
    return [l for l in linhas if any(c.strip() for c in l)]


def _texto_celula(valor) -> str:
    # Código numérico lido como 6.0 volta a ser "6"
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _achar_coluna(linhas: list[list[str]], nomes: tuple[str, ...]) -> tuple[int, int] | None:
    """(linha do cabeçalho, coluna) da primeira coluna cujo título está em `nomes`, nas 10 primeiras linhas."""

    for i, linha in enumerate(linhas[:10]):
        normalizados = [lojas_base._normalizar(c) for c in linha]
        for nome in nomes:
            if nome in normalizados:
                return i, normalizados.index(nome)
    return None


_COLUNAS_LOJA = ("FILIAL", "LOJA", "COD LOJA", "CODIGO LOJA", "COD FILIAL", "CODIGO FILIAL", "EMPRESA", "CODIGO", "COD")


def colunas_planilha(filename: str, content: bytes) -> list[dict]:
    """Colunas da planilha (cabeçalho = primeira linha preenchida), com até
    3 valores de exemplo, para o usuário escolher qual é a da loja."""

    linhas = _linhas_xlsx(filename, content)
    if not linhas:
        raise ValueError("A planilha está vazia")
    cabecalho, dados = linhas[0], linhas[1:]
    largura = max(len(l) for l in linhas)
    colunas = []
    for i in range(largura):
        nome = cabecalho[i].strip() if i < len(cabecalho) else ""
        exemplos = [l[i].strip() for l in dados if i < len(l) and l[i].strip()][:3]
        if nome or exemplos:
            colunas.append({"indice": i, "nome": nome or f"Coluna {i + 1}", "exemplos": exemplos})
    return colunas


def ler_lojas(db: Session, filename: str, content: bytes, coluna: int | None = None) -> dict:
    """Lojas de um relatório .xlsx pela coluna de código da loja. Com
    `coluna` (índice escolhido na tela), lê essa coluna abaixo do cabeçalho
    (primeira linha); sem ela, procura FILIAL, LOJA, CÓDIGO… (sem cabeçalho
    conhecido, a primeira coluna). "6" vale "06". Devolve as filiais que
    existem na base de lojas e os códigos que não casaram com nenhuma."""

    linhas = _linhas_xlsx(filename, content)
    if coluna is not None:
        cab, col = 0, coluna
    else:
        achou = _achar_coluna(linhas, _COLUNAS_LOJA)
        cab, col = achou if achou else (-1, 0)
    existentes = {l["filial"] for l in lojas_base.listar_lojas(db)}

    encontradas: list[str] = []
    nao_encontradas: list[str] = []
    for linha in linhas[cab + 1 :]:
        valor = linha[col].strip() if col < len(linha) else ""
        if not valor:
            continue
        filial = lojas_base.codigo_filial(valor)
        if filial not in existentes:
            if valor not in nao_encontradas:
                nao_encontradas.append(valor)
        elif filial not in encontradas:
            encontradas.append(filial)
    return {"lojas": encontradas, "nao_encontradas": nao_encontradas}


_COLUNAS_CODIGO = ("CODIGO", "COD", "CODIGO SETA", "COD SETA", "CODIGO CLIENTE", "COD CLIENTE", "PESSOA")
_COLUNAS_CPF = ("CPF", "CPF/CNPJ", "CPFCNPJ", "CPF CNPJ", "DOCUMENTO")


def ler_clientes(filename: str, content: bytes) -> dict:
    """Clientes de uma planilha .xlsx para restringir a campanha: usa a
    coluna de código SETA (Codigo, Código SETA…) ou, sem ela, a de CPF, que
    vira código numa consulta só ao SETA. Devolve os códigos (8 dígitos) e
    quantas linhas não deram para identificar."""

    linhas = _linhas_xlsx(filename, content)
    achou = _achar_coluna(linhas, _COLUNAS_CODIGO)
    por_cpf = False
    if achou is None:
        achou = _achar_coluna(linhas, _COLUNAS_CPF)
        por_cpf = True
    if achou is None:
        raise ValueError("Não achei a coluna de código do cliente (Codigo) nem a de CPF na planilha")
    cab, col = achou
    cabecalho = [c.strip() for c in linhas[cab]]
    colunas = [c for c in cabecalho if c]
    # (identificador só dígitos, linha como {coluna: valor})
    corpo = [
        (re.sub(r"\D", "", l[col]) if col < len(l) else "", {c: l[i] for i, c in enumerate(cabecalho) if c and i < len(l)})
        for l in linhas[cab + 1 :]
    ]
    if por_cpf:
        mapa = seta_client.codigos_por_cpf([v.zfill(11) for v, _ in corpo if v and len(v) <= 11])
        chave = lambda v: mapa.get(v.zfill(11)) if v and len(v) <= 11 else None
    else:
        chave = lambda v: v.zfill(8) if v and len(v) <= 8 else None

    por_codigo: dict[str, dict[str, str]] = {}
    ignoradas = 0
    for valor, linha in corpo:
        codigo = chave(valor)
        if codigo is None:
            ignoradas += 1
        elif codigo not in por_codigo:  # cliente repetido: vale a primeira linha
            por_codigo[codigo] = linha
    return {
        "clientes": list(por_codigo),
        "linhas": por_codigo,
        "colunas": colunas,
        "ignoradas": ignoradas,
        "coluna": "CPF" if por_cpf else "Código",
    }


def diagnosticar(db: Session, campanha: models.Campanha) -> dict:
    """Diagnóstico da campanha: pronto, aguardando_decisao, fora_por_decisao, blacklist, ja_recebeu.
    Omitimos quem não tem parcela em aberto ou está fora dos filtros (R5)."""

    todos = bool(getattr(campanha, "todos_da_planilha", False) and campanha.clientes)
    job = cobranca_base.buscar_base(
        db, **filtros_para_busca(db, campanha.filtros or {}, campanha.clientes, todos_da_planilha=todos)
    )
    if job["status"] != "ready":
        return {"status": "processing"}
        
    regras = carregar_regras(db)
    if todos:
        base = [
            {**c, "faixa": faixa}
            for c in job["data"]
            if (faixa := faixa_na_campanha(regras, c["dias_atraso"])) is not None
        ]
    else:
        so_campanhas = regras.nomes_faixa_so_campanhas
        base = [c for c in job["data"] if c["dias_atraso"] >= 1 or c["faixa"] in so_campanhas]
        
    ja_foram = ja_receberam(db, campanha)
    bloqueados = _bloqueados(db, campanha)
    
    def status_variaveis(codigo: str) -> str:
        linha = (campanha.planilha_linhas or {}).get(codigo)
        if not linha:
            return "pronto"
        if campanha.fonte_valores == "planilha" and valor_da_planilha(linha) is None:
            # Valor inválido na planilha: o do SETA só com autorização do usuário
            if campanha.valor_seta_autorizado is None:
                return "aguardando_decisao"
            if not campanha.valor_seta_autorizado:
                return "fora_por_decisao"

        for m in campanha.faixa.variable_mappings:
            if m.fonte_tipo == "coluna" and m.column_name:
                val = str(linha.get(m.column_name, "")).strip()
                if not val:
                    if m.reserva_vazio == "fora":
                        return "fora_por_decisao"
                    elif not m.reserva_vazio:
                        return "aguardando_decisao"
        return "pronto"

    diagnostico = {
        "pronto": 0,
        "aguardando_decisao": 0,
        "fora_por_decisao": 0,
        "blacklist": 0,
        "ja_recebeu": 0
    }
    
    clientes_na_previa = []
    
    for c in base:
        if campanha.clientes and c["codigo"] not in campanha.planilha_linhas:
            continue
            
        codigo = c["codigo"]
        if codigo in ja_foram:
            diagnostico["ja_recebeu"] += 1
        elif codigo in bloqueados:
            diagnostico["blacklist"] += 1
        else:
            st = status_variaveis(codigo)
            diagnostico[st] += 1
            if st == "pronto":
                clientes_na_previa.append(c)
                
    return {
        "status": "ready",
        "total_base": len(base),
        "clientes": clientes_na_previa,
        "diagnostico": diagnostico
    }
