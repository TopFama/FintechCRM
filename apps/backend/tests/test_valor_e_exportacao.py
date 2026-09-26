"""Validação da matriz de valor em aberto e da exportação de leads para .xlsx.

Executa com assert simples, sem pytest. Encerra imprimindo 'OK'.
"""

from datetime import date, datetime
from decimal import Decimal
import io
import os
import tempfile
from unittest.mock import patch

# Variáveis mínimas para inicializar a aplicação contra o banco agy_e
os.environ["DATABASE_URL"] = "postgresql+psycopg://postgres:t@localhost:15432/agy_e"
os.environ["JWT_SECRET"] = "segredo-de-teste-longo-e-seguro-12345"
os.environ["ADMIN_PASSWORD"] = "senha-admin-teste"
os.environ["MEDIA_DIR"] = tempfile.gettempdir()

from fastapi.testclient import TestClient
import openpyxl

import os as _os
from cryptography.fernet import Fernet as _Fernet

_os.environ.setdefault("ENCRYPTION_KEY", _Fernet.generate_key().decode())  # o backend não sobe sem ela
from app import models
from app.cobranca_regras import Cluster, FaixaAtraso, ParametrosJuros, Regras
from app.cobranca_relatorio import montar_matriz, montar_matriz_valor
from app.database import SessionLocal
from app.main import app
from app.regras_db import carregar_regras


# ---------------------------------------------------------------------------
# 1. Testes puros: montar_matriz_valor e montar_matriz
# ---------------------------------------------------------------------------

regras_puras = Regras(
    clusters=(
        Cluster(nome="ESPECIAL", valor_min=Decimal("0")),
        Cluster(nome="POTENCIAL", valor_min=Decimal("1000")),
    ),
    faixas=(
        FaixaAtraso(nome="11 A 20", dia_min=11, dia_max=20),
        FaixaAtraso(nome="21 A 30", dia_min=21, dia_max=30),
    ),
    whatsapp=frozenset([("ESPECIAL", "11 A 20"), ("POTENCIAL", "21 A 30")]),
    juros=ParametrosJuros(),
)

# 4 clientes simulados em 2 clusters / 2 faixas (um com faixa=None ignorado)
clientes_fake = [
    {
        "codigo": "00000001",
        "cluster": "ESPECIAL",
        "faixa": "11 A 20",
        "valor_em_aberto": Decimal("150.00"),
        "spc_restricao": "nao",
    },
    {
        "codigo": "00000002",
        "cluster": "ESPECIAL",
        "faixa": "21 A 30",
        "valor_em_aberto": Decimal("250.50"),
        "spc_restricao": "sim",
    },
    {
        "codigo": "00000003",
        "cluster": "POTENCIAL",
        "faixa": "21 A 30",
        "valor_em_aberto": Decimal("300.25"),
        "spc_restricao": "sim",
    },
    {
        "codigo": "00000004",
        "cluster": "POTENCIAL",
        "faixa": None,  # ignorado
        "valor_em_aberto": Decimal("500.00"),
        "spc_restricao": "nao",
    },
]

# Resultado de montar_matriz (quantidade inalterada)
matriz_qtd = montar_matriz(clientes_fake, regras_puras)
assert matriz_qtd["celulas"] == {
    "ESPECIAL": {"11 A 20": 1, "21 A 30": 1},
    "POTENCIAL": {"11 A 20": 0, "21 A 30": 1},
}
assert matriz_qtd["total_por_cluster"] == {"ESPECIAL": 2, "POTENCIAL": 1}
assert matriz_qtd["total_por_faixa"] == {"11 A 20": 1, "21 A 30": 2}
assert matriz_qtd["total"] == 3

# Resultado de montar_matriz_valor (valores exatos em Decimal, células vazias Decimal("0"))
matriz_val = montar_matriz_valor(clientes_fake, regras_puras)
assert matriz_val["celulas"] == {
    "ESPECIAL": {"11 A 20": Decimal("150.00"), "21 A 30": Decimal("250.50")},
    "POTENCIAL": {"11 A 20": Decimal("0"), "21 A 30": Decimal("300.25")},
}
assert matriz_val["total_por_cluster"] == {
    "ESPECIAL": Decimal("400.50"),
    "POTENCIAL": Decimal("300.25"),
}
assert matriz_val["total_por_faixa"] == {
    "11 A 20": Decimal("150.00"),
    "21 A 30": Decimal("550.75"),
}
assert matriz_val["total"] == Decimal("700.75")

# Verifica se todos os tipos numéricos são Decimal
for c in regras_puras.nomes_cluster:
    for f in regras_puras.nomes_faixa:
        assert isinstance(matriz_val["celulas"][c][f], Decimal)
for v in matriz_val["total_por_cluster"].values():
    assert isinstance(v, Decimal)
for v in matriz_val["total_por_faixa"].values():
    assert isinstance(v, Decimal)
assert isinstance(matriz_val["total"], Decimal)


# ---------------------------------------------------------------------------
# 2. Testes de API com TestClient (Postgres agy_e)
# ---------------------------------------------------------------------------

ADMIN_PASSWORD = "senha-admin-teste"
os.environ["ADMIN_PASSWORD"] = ADMIN_PASSWORD

from app.security import hash_password

with TestClient(app) as client:
    # Garante que o admin tenha a senha esperada pelo teste
    db_init = SessionLocal()
    admin_user = db_init.query(models.User).filter(models.User.email == "admin@topfama.com.br").first()
    if admin_user:
        admin_user.password_hash = hash_password(ADMIN_PASSWORD)
    else:
        db_init.add(models.User(email="admin@topfama.com.br", password_hash=hash_password(ADMIN_PASSWORD)))
    db_init.commit()
    db_init.close()

    # Autenticação
    res_login = client.post(
        "/auth/login",
        json={"email": "admin@topfama.com.br", "password": ADMIN_PASSWORD},
    )
    assert res_login.status_code == 200, res_login.text
    token = res_login.json()["access_token"]
    auth_headers = {"Authorization": f"Bearer {token}"}

    db = SessionLocal()
    regras_db = carregar_regras(db)
    cluster_1 = regras_db.nomes_cluster[0]
    cluster_2 = regras_db.nomes_cluster[1]
    faixa_1 = regras_db.nomes_faixa[3]  # ex: "11 A 20"
    faixa_2 = regras_db.nomes_faixa[4]  # ex: "21 A 30"

    # API /cobranca/relatorio com cobranca_base.buscar_base mockado
    fake_clientes_relatorio = [
        {
            "codigo": "00000001",
            "cluster": cluster_1,
            "faixa": faixa_1,
            "valor_em_aberto": Decimal("100.50"),
            "spc_restricao": "nao",
        },
        {
            "codigo": "00000002",
            "cluster": cluster_1,
            "faixa": faixa_2,
            "valor_em_aberto": Decimal("200.25"),
            "spc_restricao": "sim",
        },
        {
            "codigo": "00000003",
            "cluster": cluster_2,
            "faixa": faixa_2,
            "valor_em_aberto": Decimal("300.00"),
            "spc_restricao": "nao",
        },
    ]

    with patch("app.routers.cobranca.cobranca_base.buscar_base", return_value={"status": "ready", "data": fake_clientes_relatorio}):
        res_rel = client.get("/cobranca/relatorio", headers=auth_headers)
        assert res_rel.status_code == 200, res_rel.text
        rel_data = res_rel.json()["data"]  # resposta assíncrona: {"status", "data"}

        assert "valor_em_aberto" in rel_data
        assert "quantidade" in rel_data

        # Quantidade inalterada
        qtd = rel_data["quantidade"]
        assert qtd["celulas"][cluster_1][faixa_1] == 1
        assert qtd["celulas"][cluster_1][faixa_2] == 1
        assert qtd["celulas"][cluster_2][faixa_2] == 1
        assert qtd["total"] == 3

        # Valor em aberto serializado em Decimal string
        val = rel_data["valor_em_aberto"]
        assert val["celulas"][cluster_1][faixa_1] == "100.50"
        assert val["celulas"][cluster_1][faixa_2] == "200.25"
        assert val["celulas"][cluster_2][faixa_2] == "300.00"
        assert val["total_por_cluster"][cluster_1] == "300.75"
        assert val["total_por_cluster"][cluster_2] == "300.00"
        assert val["total_por_faixa"][faixa_1] == "100.50"
        assert val["total_por_faixa"][faixa_2] == "500.25"
        assert val["total"] == "600.75"


    # -----------------------------------------------------------------------
    # 3. Exportação de leads para .xlsx (/leads/exportar.xlsx)
    # -----------------------------------------------------------------------

    # Limpeza da base de teste
    db.query(models.Lead).delete()
    db.query(models.ClienteBloqueado).delete()
    db.commit()

    f_antiga = faixa_1  # anterior na ordem configurada
    f_recente = faixa_2  # posterior na ordem configurada
    hoje = date.today().isoformat()

    # Lead 1: cobrado, faixa f_recente
    lead1 = models.Lead(
        codigo_cliente="123456",
        nome="MARIA DA SILVA",
        cpf="52998224725",
        celular="63991234567",
        cluster=cluster_1,
        faixa=f_recente,
        dias_atraso=25,
        valor_em_aberto=Decimal("150.00"),
        valor_cobrar=Decimal("160.00"),
        vencimento_mais_antigo=date(2026, 1, 1),
        status="cobrado",
    )
    # Lead 2: cobrado, faixa f_antiga, CPF incompleto e sem celular
    lead2 = models.Lead(
        codigo_cliente="456",
        nome="PEDRO ALVES",
        cpf="998224725",  # 9 dígitos: formatar_cpf completa com zeros à esquerda
        celular=None,     # sem celular
        cluster=cluster_1,
        faixa=f_antiga,
        dias_atraso=15,
        valor_em_aberto=Decimal("200.00"),
        valor_cobrar=Decimal("210.00"),
        vencimento_mais_antigo=date(2026, 1, 10),
        status="cobrado",
    )
    # Lead 2b: cobrado, faixa f_antiga, nome anterior a "Pedro" (testa ordenação secundária por nome)
    lead2b = models.Lead(
        codigo_cliente="222",
        nome="BRUNO TESTE",
        cpf="12345678901",
        celular="63999998888",
        cluster=cluster_1,
        faixa=f_antiga,
        dias_atraso=14,
        valor_em_aberto=Decimal("180.00"),
        valor_cobrar=Decimal("190.00"),
        vencimento_mais_antigo=date(2026, 1, 11),
        status="cobrado",
    )
    # Lead 3: novo, faixa f_antiga
    lead3 = models.Lead(
        codigo_cliente="789",
        nome="ANA SOUZA",
        cpf="11222333000181",
        celular="(63) 99123-4567",
        cluster=cluster_2,
        faixa=f_antiga,
        dias_atraso=12,
        valor_em_aberto=Decimal("300.00"),
        valor_cobrar=Decimal("310.00"),
        vencimento_mais_antigo=date(2026, 1, 15),
        status="novo",
    )
    # Lead 4: cobrado, mas cliente está na blacklist (deve ser excluído)
    lead4 = models.Lead(
        codigo_cliente="99999999",
        nome="CARLOS BLOQUEADO",
        cpf="11122233344",
        celular="63991234567",
        cluster=cluster_1,
        faixa=f_recente,
        dias_atraso=22,
        valor_em_aberto=Decimal("400.00"),
        valor_cobrar=Decimal("410.00"),
        vencimento_mais_antigo=date(2026, 1, 5),
        status="cobrado",
    )
    bloqueado = models.ClienteBloqueado(
        tipo="seta",
        valor="99999999",
        motivo="Bloqueio teste blacklist",
    )

    db.add_all([lead1, lead2, lead2b, lead3, lead4, bloqueado])
    db.commit()

    # 1. 401 sem login
    res_sem_login = client.get("/leads/exportar.xlsx")
    assert res_sem_login.status_code == 401, f"esperava 401, obteve {res_sem_login.status_code}"

    # 2. Chamada padrão (retorna apenas status="cobrado", sem blacklist)
    res_padrao = client.get("/leads/exportar.xlsx", headers=auth_headers)
    assert res_padrao.status_code == 200, res_padrao.text
    assert res_padrao.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert res_padrao.headers["content-disposition"] == f'attachment; filename="leads_todas_{hoje}.xlsx"'

    wb_padrao = openpyxl.load_workbook(io.BytesIO(res_padrao.content))
    ws_padrao = wb_padrao["Leads"]

    # Cabeçalho exato
    assert [ws_padrao.cell(1, c).value for c in range(1, 5)] == ["Codigo", "Nome", "Celular", "CPF"]

    # Total de 4 linhas: 1 cabeçalho + 3 dados (lead2b, lead2, lead1).
    # lead3 (novo) e lead4 (blacklist) foram excluídos.
    # Ordem de faixa: f_antiga vem antes de f_recente.
    # Ordem de nome dentro de f_antiga: Bruno vem antes de Pedro.
    assert ws_padrao.max_row == 4, f"esperava 4 linhas, obteve {ws_padrao.max_row}"

    linha2 = [ws_padrao.cell(2, c).value for c in range(1, 5)]
    assert linha2[0] == "00000222"
    assert linha2[1] == "Bruno"
    assert linha2[3] == "123.456.789-01"
    assert linha2[2] == "5563999998888"

    linha3 = [ws_padrao.cell(3, c).value for c in range(1, 5)]
    assert linha3[0] == "00000456"
    assert linha3[1] == "Pedro"
    assert linha3[3] == "009.982.247-25"
    assert linha3[2] in ("", None)

    linha4 = [ws_padrao.cell(4, c).value for c in range(1, 5)]
    assert linha4[0] == "00123456"
    assert linha4[1] == "Maria"
    assert linha4[3] == "529.982.247-25"
    assert linha4[2] == "5563991234567"

    # 3. ?faixa= restringe a faixa
    res_faixa = client.get(f"/leads/exportar.xlsx?faixa={f_recente}", headers=auth_headers)
    assert res_faixa.status_code == 200
    nome_faixa_limpo = f_recente.replace(" ", "")
    assert res_faixa.headers["content-disposition"] == f'attachment; filename="leads_{nome_faixa_limpo}_{hoje}.xlsx"'
    ws_faixa = openpyxl.load_workbook(io.BytesIO(res_faixa.content))["Leads"]
    assert ws_faixa.max_row == 2
    assert ws_faixa.cell(2, 1).value == "00123456"
    assert ws_faixa.cell(2, 2).value == "Maria"

    # 4. ?status=novo funciona
    res_novo = client.get("/leads/exportar.xlsx?status=novo", headers=auth_headers)
    assert res_novo.status_code == 200
    ws_novo = openpyxl.load_workbook(io.BytesIO(res_novo.content))["Leads"]
    assert ws_novo.max_row == 2
    assert ws_novo.cell(2, 1).value == "00000789"
    assert ws_novo.cell(2, 2).value == "Ana"
    assert ws_novo.cell(2, 3).value == "5563991234567"
    assert ws_novo.cell(2, 4).value == "11.222.333/0001-81"  # colunas: Codigo, Nome, Celular, CPF

    # 5. Resultado vazio -> planilha válida somente com cabeçalho (não 404)
    res_vazio = client.get("/leads/exportar.xlsx?faixa=151%2B", headers=auth_headers)
    assert res_vazio.status_code == 200
    assert res_vazio.headers["content-disposition"] == f'attachment; filename="leads_151+_{hoje}.xlsx"'
    ws_vazio = openpyxl.load_workbook(io.BytesIO(res_vazio.content))["Leads"]
    assert ws_vazio.max_row == 1
    assert [ws_vazio.cell(1, c).value for c in range(1, 5)] == ["Codigo", "Nome", "Celular", "CPF"]

    db.close()

print("OK")
