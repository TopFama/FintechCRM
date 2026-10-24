"""Filtros por coluna (estilo Excel) dos relatórios: pendentes, erros, envios,
telefones inválidos e quem pagou. Cada relatório filtra, lista os valores da
coluna (dentro dos demais filtros) e exporta com o mesmo filtro.

Executa com assert simples, sem pytest. Encerra imprimindo 'OK'.
"""

import io
import json
import os
import tempfile
from datetime import datetime, timedelta
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import patch

from cryptography.fernet import Fernet

os.environ["DATABASE_URL"] = "postgresql+psycopg://postgres:t@localhost:15432/agy_filtros"
os.environ["JWT_SECRET"] = "segredo-de-teste-filtros-coluna-123"
os.environ["ADMIN_PASSWORD"] = "senha-admin-teste-filtros"
os.environ["MEDIA_DIR"] = tempfile.mkdtemp()
os.environ.setdefault("ENCRYPTION_KEY", Fernet.generate_key().decode())

import openpyxl
from fastapi.testclient import TestClient

from app import models
from app.database import SessionLocal
from app.main import app

with TestClient(app):  # sobe o schema e cria o admin
    pass

client = TestClient(app)
res = client.post("/auth/login", json={"email": "admin@topfama.com.br", "password": os.environ["ADMIN_PASSWORD"]})
assert res.status_code == 200, res.text
auth = {"Authorization": f"Bearer {res.json()['access_token']}"}

db = SessionLocal()
S = models.QueueStatus
agora = datetime.utcnow()
regua, regua2 = models.Faixa(name="1 A 10"), models.Faixa(name="11 A 20")
campanha = models.Faixa(name="Campanha: Black Friday", tipo="campanha")
numero = models.WhatsappNumber(phone_number_id="pn-filtros", display_phone_number="5563999990001", waba_id="w1")
db.add_all([regua, regua2, campanha, numero])
db.flush()


def item(faixa, codigo, status, lojas=",07,", valor=None, faixa_atraso=None, nome="Ana", celular="5563911110000",
         erro=None, criado=agora, enviado=None):
    return models.QueueItem(
        faixa_id=faixa.id, codigo_cliente=codigo, nome=nome, celular=celular, celular_original=celular, status=status,
        lojas=lojas, valor=valor, faixa_atraso=faixa_atraso, error_message=erro, created_at=criado,
        sent_at=enviado, whatsapp_number_id=numero.id if status == S.sent else None,
    )


db.add_all([
    item(regua, "P1", S.pending, lojas=",07,", valor="1.500,00"),
    item(regua2, "P2", S.pending, lojas=",07,09,", valor="300.00"),
    item(campanha, "P3", S.pending, lojas=",", valor="300.00", faixa_atraso="1 A 10"),
    item(regua, "E1", S.error, erro="Template recusado"),
    item(regua2, "E2", S.error, erro=None, criado=agora - timedelta(days=10)),
    item(regua, "S1", S.sent, enviado=agora, nome="Bia"),
    item(regua2, "S2", S.sent, enviado=agora, nome="Caio"),
    item(campanha, "S3", S.sent, enviado=agora, nome="Caio", faixa_atraso="1 A 10"),
])
db.add_all([
    models.InvalidPhoneRecord(
        faixa_id=regua.id, codigo_cliente="I1", celular_original="6399", celular_normalizado="556399",
        motivo="curto", lojas=",07,", created_at=agora,
    ),
    models.InvalidPhoneRecord(
        faixa_id=regua.id, codigo_cliente="I2", celular_original="6388", celular_normalizado=None,
        motivo="letras", lojas=",", created_at=agora,
    ),
])
db.commit()


def f(**colunas):
    return {"filtros": json.dumps(colunas)}


def get(url, **params):
    r = client.get(url, params=params, headers=auth)
    assert r.status_code == 200, (url, params, r.text)
    return r.json()


def codigos(url, **params):
    return sorted(i["codigo_cliente"] for i in get(url, **params)["itens"])


def valores(url, coluna, busca=None, **params):
    r = get(url, coluna=coluna, **({"busca": busca} if busca else {}), **params)
    return r["valores"]


# --- Pendentes -----------------------------------------------------------------
P = "/relatorios/pendentes"
assert codigos(P) == ["P1", "P2", "P3"]
# Faixa de atraso: a campanha cai na faixa de atraso do cliente (como a coluna da tela)
assert codigos(P, **f(faixa={"valores": ["1 A 10"]})) == ["P1", "P3"]
assert codigos(P, **f(faixa={"valores": []})) == []
# Campanha vazia = régua
assert codigos(P, **f(campanha={"valores": [""]})) == ["P1", "P2"]
assert codigos(P, **f(campanha={"valores": ["Black Friday"]})) == ["P3"]
# Loja: item com várias lojas aparece em cada uma; "," = sem loja
assert codigos(P, **f(loja={"valores": ["09"]})) == ["P2"]
assert codigos(P, **f(loja={"valores": ["07"]})) == ["P1", "P2"]
assert codigos(P, **f(loja={"valores": [""]})) == ["P3"]
# Valor: lista pelo número (1.500,00 e 1500.00 são o mesmo valor) e intervalo
assert codigos(P, **f(valor={"valores": ["300.00"]})) == ["P2", "P3"]
assert codigos(P, **f(valor={"de": "1000"})) == ["P1"]
assert codigos(P, **f(valor={"de": "100", "ate": "300"})) == ["P2", "P3"]
# Filtros se somam e valem na contagem e na paginação
pagina = get(P, limit=1, offset=1, **f(loja={"valores": ["07"]}))
assert pagina["total"] == 2 and len(pagina["itens"]) == 1
# Lista de valores: dentro dos outros filtros, sem o da própria coluna
assert valores(P + "/valores", "loja") == ["", "07", "09"]
assert valores(P + "/valores", "loja", **f(valor={"valores": ["300.00"]})) == ["", "07", "09"]
assert valores(P + "/valores", "loja", **f(valor={"valores": ["1500.00"]})) == ["07"]
assert valores(P + "/valores", "loja", **f(loja={"valores": ["09"]}, valor={"valores": ["1500.00"]})) == ["07"]
assert valores(P + "/valores", "campanha") == ["", "Black Friday"]
assert valores(P + "/valores", "campanha", "black") == ["Black Friday"]
assert [Decimal(v) for v in valores(P + "/valores", "valor")] == [Decimal("300"), Decimal("1500")]
# Excel com o mesmo filtro da tela
xlsx = client.get(P + "/export", params=f(loja={"valores": ["09"]}), headers=auth)
linhas = list(openpyxl.load_workbook(io.BytesIO(xlsx.content)).active.iter_rows(values_only=True))
assert [l[0] for l in linhas[1:]] == ["P2"], linhas
# Descartar fila usa os mesmos filtros da tela
previa = get(P + "/descartar/previa", **f(loja={"valores": ["09"]}))
assert previa["qtd"] == 1

# --- Erros ---------------------------------------------------------------------
E = "/relatorios/erros"
assert codigos(E) == ["E1", "E2"]
assert codigos(E, **f(mensagem={"valores": ["Template recusado"]})) == ["E1"]
assert codigos(E, **f(mensagem={"valores": ["Erro sem detalhe"]})) == ["E2"]
# Quando: intervalo de datas (de, até)
hoje = (agora - timedelta(hours=3)).date().isoformat()
assert codigos(E, **f(quando={"de": hoje})) == ["E1"]
assert codigos(E, **f(quando={"ate": (agora - timedelta(days=5)).date().isoformat()})) == ["E2"]
assert valores(E + "/valores", "mensagem") == ["Erro sem detalhe", "Template recusado"]
assert client.get(E + "/valores", params={"coluna": "quando"}, headers=auth).status_code == 422

# --- Envios realizados -----------------------------------------------------------
V = "/relatorios/envios"
assert codigos(V) == ["S1", "S2", "S3"]
assert codigos(V, **f(faixa={"valores": ["11 A 20"]})) == ["S2"]
# faixa de atraso traz também a campanha feita a quem estava nela (como o "Por faixa" do Dashboard)
assert codigos(V, **f(faixa={"valores": ["1 A 10"]})) == ["S1", "S3"]
assert codigos(V, **f(faixa={"valores": ["Campanha: Black Friday"]})) == ["S3"]
assert codigos(V, **f(nome={"valores": ["Caio"]}, faixa={"valores": ["11 A 20", "Campanha: Black Friday"]})) == ["S2", "S3"]
assert codigos(V, **f(telefone={"valores": ["5563999990001"]})) == ["S1", "S2", "S3"]
assert codigos(V, **f(telefone={"valores": ["000"]})) == []
assert valores(V + "/valores", "faixa") == ["1 A 10", "11 A 20", "Campanha: Black Friday"]
assert valores(V + "/valores", "faixa", **f(nome={"valores": ["Bia"]})) == ["1 A 10"]
assert valores(V + "/valores", "nome", "ai") == ["Caio"]
xlsx = client.get(V + "/export", params=f(nome={"valores": ["Bia"]}), headers=auth)
linhas = list(openpyxl.load_workbook(io.BytesIO(xlsx.content)).active.iter_rows(values_only=True))
assert [l[0] for l in linhas[1:]] == ["S1"], linhas
# o telefone cobrado sai como a tela mostra, no Excel; a pesquisa da lista aceita o formato exibido
assert linhas[1][linhas[0].index("Telefone cobrado")] == "+55 (63) 91111-0000", linhas
assert valores(V + "/valores", "celular", "(63) 91111") == ["5563911110000"]
# a ordenação continua valendo junto com o filtro
assert get(V, sort_by="faixa", sort_dir="desc", **f(nome={"valores": ["Caio"]}))["total"] == 2

# --- Telefones inválidos ---------------------------------------------------------
I = "/relatorios/telefones-invalidos"
assert codigos(I) == ["I1", "I2"]
assert codigos(I, **f(motivo={"valores": ["curto"]})) == ["I1"]
assert codigos(I, **f(celular_normalizado={"valores": [""]})) == ["I2"]
assert codigos(I, **f(loja={"valores": ["07"]})) == ["I1"]
assert valores(I + "/valores", "motivo") == ["curto", "letras"]
assert valores(I + "/valores", "motivo", **f(loja={"valores": ["07"]})) == ["curto"]
xlsx = client.get(I + "/export", params=f(motivo={"valores": ["letras"]}), headers=auth)
linhas = list(openpyxl.load_workbook(io.BytesIO(xlsx.content)).active.iter_rows(values_only=True))
assert [l[0] for l in linhas[1:]] == ["I2"], linhas

# --- Quem pagou (lista em memória do cache do SETA) --------------------------------
from datetime import date

linhas_pagou = [
    {"codigo_cliente": "C1", "nome": "Ana", "cpf": None, "loja": "LOJA 07, LOJA 09", "faixa": "1 A 10, 11 A 20",
     "data_cobranca": date.today(), "valor_cobrado": Decimal("100.00"), "valor_pago": Decimal("100.00"),
     "qtd_titulos_pagos": 1, "primeiro_pagamento": date.today(), "ultimo_pagamento": date.today()},
    {"codigo_cliente": "C2", "nome": "", "cpf": None, "loja": "", "faixa": "11 A 20",
     "data_cobranca": date.today(), "valor_cobrado": Decimal("900.00"), "valor_pago": Decimal("50.00"),
     "qtd_titulos_pagos": 1, "primeiro_pagamento": date.today(), "ultimo_pagamento": date.today()},
]
snap = SimpleNamespace(data=linhas_pagou, gerado_em=datetime.utcnow().timestamp(), velho=False)
PG = "/relatorios/pagamentos"
with patch("app.routers.reports._pagamentos", return_value=snap):
    assert codigos(PG) == ["C1", "C2"]
    assert codigos(PG, **f(loja={"valores": ["LOJA 09"]})) == ["C1"]
    assert codigos(PG, **f(loja={"valores": [""]})) == ["C2"]
    assert codigos(PG, **f(nome={"valores": [""]})) == ["C2"]
    assert codigos(PG, **f(valor_cobrado={"de": "500"})) == ["C2"]
    assert codigos(PG, **f(valor_pago={"valores": ["50.00"]})) == ["C2"]
    resumo = get(PG, **f(loja={"valores": ["LOJA 07"]}))
    assert resumo["total"] == 1 and Decimal(str(resumo["valor_pago"])) == Decimal("100.00")
    assert valores(PG + "/valores", "faixa") == ["1 A 10", "11 A 20"]
    assert valores(PG + "/valores", "loja", **f(valor_cobrado={"de": "500"})) == [""]
    assert valores(PG + "/valores", "valor_cobrado") == ["100.00", "900.00"]
    xlsx = client.get(PG + "/export", params=f(valor_pago={"valores": ["50.00"]}), headers=auth)
    linhas = list(openpyxl.load_workbook(io.BytesIO(xlsx.content)).active.iter_rows(values_only=True))
    assert [l[0] for l in linhas[1:]] == ["C2"], linhas

# --- Clientes da Cobrança (lista em memória da base do SETA) -----------------------
def cliente(codigo, nome, faixa, valor, dias, venc):
    return {"codigo": codigo, "nome": nome, "celular": "5563911110000", "celular_origem": "celular1",
            "celular_original": "63911110000", "cpfcnpj": "1" * 11, "status": "A", "status_descricao": "Ativo",
            "loja_cadastro": None, "salario": None, "limite_rotativo": None, "nascimento": None, "cadastro": None,
            "cluster": "ESPECIAL", "valor_pago": Decimal("0"), "qtd_compras": 1, "faixa_compra": None,
            "ultima_compra": None, "faixa": faixa, "dias_atraso": dias, "entra_whatsapp": True, "qtd_titulos": 1,
            "valor_em_aberto": Decimal(valor), "qtd_parcelas_cobranca": 1, "valor_cobrar": Decimal(valor),
            "vencimento_mais_antigo": venc, "lojas": [], "portadores": [], "spc_restricao": "nao"}


base_cobranca = [
    cliente("K1", "Ana", "1 A 10", "100.00", 5, date(2026, 10, 1)),
    cliente("K2", "Bia", "11 A 20", "900.00", 15, date(2026, 9, 20)),
    cliente("K3", "Caio", None, "50.00", 1, date(2026, 10, 10)),
]
CL = "/cobranca/clientes"
with patch("app.routers.cobranca.buscar_base_ou_erro", return_value={"status": "ready", "data": base_cobranca}), \
        patch("app.routers.cobranca.seta_client.buscar_spc", return_value={}):
    def clientes(**params):
        r = get(CL, **params)["data"]
        return sorted(c["codigo"] for c in r["itens"]), r["total"]

    assert clientes() == (["K1", "K2", "K3"], 3)
    assert clientes(**f(faixa={"valores": ["1 A 10"]})) == (["K1"], 1)
    assert clientes(**f(faixa={"valores": [""]})) == (["K3"], 1)
    assert clientes(**f(valor_cobrar={"de": "100"})) == (["K1", "K2"], 2)
    assert clientes(**f(dias_atraso={"ate": "5"})) == (["K1", "K3"], 2)
    assert clientes(**f(vencimento_mais_antigo={"de": "2026-10-01"})) == (["K1", "K3"], 2)
    assert clientes(**f(vencimento_mais_antigo={"ate": "2026-09-30"})) == (["K2"], 1)
    assert clientes(**f(nome={"valores": ["Ana", "Bia"]}, dias_atraso={"de": "10"})) == (["K2"], 1)
    assert valores(CL + "/valores", "faixa") == ["", "1 A 10", "11 A 20"]
    assert valores(CL + "/valores", "nome", **f(valor_cobrar={"de": "100"})) == ["Ana", "Bia"]
    assert valores(CL + "/valores", "nome", "ca") == ["Caio"]
    # data filtra por intervalo: não tem lista
    assert client.get(CL + "/valores", params={"coluna": "vencimento_mais_antigo"}, headers=auth).status_code == 422
    xlsx = client.get(CL + "/exportar.xlsx", params=f(faixa={"valores": ["11 A 20"]}), headers=auth)
    linhas = list(openpyxl.load_workbook(io.BytesIO(xlsx.content)).active.iter_rows(values_only=True))
    assert [l[0] for l in linhas[1:]] == ["K2"], linhas
    assert client.get(CL, params=f(inexistente={"valores": ["x"]}), headers=auth).status_code == 422

# --- Entrada inválida --------------------------------------------------------------
for url in (P, E, V, I, PG):
    assert client.get(url, params={"filtros": "{"}, headers=auth).status_code == 422
    assert client.get(url, params=f(inexistente={"valores": ["x"]}), headers=auth).status_code == 422
assert client.get(P, params=f(valor={"de": "abc"}), headers=auth).status_code == 422
assert client.get(P, params=f(entrou_em={"de": "ontem"}), headers=auth).status_code == 422

db.close()
print("OK")
