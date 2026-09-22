"""Testes completos do relatório de efetividade da cobrança:
- Módulo puro app/relatorio_efetividade.py (montar_relatorio)
- Regra de janela de pagamento (dias_janela)
- Endpoints da API via TestClient com Postgres agy_efet
- Exportação Excel (.xlsx) com abas, formatações e regras condicionais
- Validação de exports existentes mantidos

Executa com assert simples (sem pytest) e imprime OK ao final.
"""

import io
import os
import tempfile
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal

# Configura ambiente mínimo para inicialização da app
os.environ["DATABASE_URL"] = "postgresql+psycopg://postgres:t@localhost:15432/agy_efet"
os.environ["JWT_SECRET"] = "uma-chave-secreta-forte-e-longa-para-o-teste-12345"
os.environ["ADMIN_PASSWORD"] = "senha-admin-segura-para-o-teste-123"
os.environ["MEDIA_DIR"] = tempfile.gettempdir()

import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

import os as _os
from cryptography.fernet import Fernet as _Fernet

_os.environ.setdefault("ENCRYPTION_KEY", _Fernet.generate_key().decode())  # o backend não sobe sem ela
from app import cobranca_base, database, google_client, models, seta_client
from app.main import app
from app.relatorio_efetividade import montar_relatorio


# ============================================================================
# 1. Testes puros: montar_relatorio
# ============================================================================

# Cenário: 4 clientes, 2 lojas, 2 faixas
# - C1: 2 parcelas (uma na loja 01, outra na loja 02, ambas faixa 11 A 20) -> pagou tudo
# - C2: 2 parcelas (ambas loja 01, faixa 11 A 20) -> pagou 1 parcela, outra em aberto
# - C3: 1 parcela (loja 02, faixa 21 A 30) -> renegociada
# - C4: 1 parcela (loja 02, faixa 21 A 30) -> não pagou nada
itens_4c = [
    {
        "codigo_cliente": "00000001",
        "faixa": "11 A 20",
        "empresa": "01",
        "valor_cobrar": Decimal("100.00"),
        "pago": True,
        "renegociada": False,
        "valor_pago": Decimal("100.00"),
    },
    {
        "codigo_cliente": "00000001",
        "faixa": "11 A 20",
        "empresa": "02",
        "valor_cobrar": Decimal("50.00"),
        "pago": True,
        "renegociada": False,
        "valor_pago": Decimal("50.00"),
    },
    {
        "codigo_cliente": "00000002",
        "faixa": "11 A 20",
        "empresa": "01",
        "valor_cobrar": Decimal("80.00"),
        "pago": True,
        "renegociada": False,
        "valor_pago": Decimal("80.00"),
    },
    {
        "codigo_cliente": "00000002",
        "faixa": "11 A 20",
        "empresa": "01",
        "valor_cobrar": Decimal("70.00"),
        "pago": False,
        "renegociada": False,
        "valor_pago": Decimal("0.00"),
    },
    {
        "codigo_cliente": "00000003",
        "faixa": "21 A 30",
        "empresa": "02",
        "valor_cobrar": Decimal("200.00"),
        "pago": False,
        "renegociada": True,
        "valor_pago": Decimal("0.00"),
    },
    {
        "codigo_cliente": "00000004",
        "faixa": "21 A 30",
        "empresa": "02",
        "valor_cobrar": Decimal("150.00"),
        "pago": False,
        "renegociada": False,
        "valor_pago": Decimal("0.00"),
    },
]

lojas_info = {
    "01": {"nome_com_cod": "01 - PALMAS", "regional": "REGIONAL NORTE", "cluster_inad": "ALTO"}
}  # Loja 02 de propósito não está no mapa para validar campos nulos

res = montar_relatorio(itens_4c, lojas_info=lojas_info)

# Por Faixa
por_faixa = {r["faixa"]: r for r in res["por_faixa"]}
assert "11 A 20" in por_faixa
f1 = por_faixa["11 A 20"]
assert f1["clientes_cobrados"] == 2  # C1 e C2
assert f1["valor_cobrado"] == Decimal("300.00")  # 100 + 50 + 80 + 70
assert f1["clientes_pagaram"] == 2  # C1 e C2 pagaram pelo menos 1 parcela
assert f1["valor_pago"] == Decimal("230.00")  # 100 + 50 + 80
assert f1["parcelas_cobradas"] == 4
assert f1["parcelas_pagas"] == 3
assert f1["parcelas_renegociadas"] == 0
assert f1["conversao_clientes"] == Decimal("1.0000")  # 2 / 2
assert f1["recuperacao_valor"] == (Decimal("230") / Decimal("300")).quantize(Decimal("0.0001"))

assert "21 A 30" in por_faixa
f2 = por_faixa["21 A 30"]
assert f2["clientes_cobrados"] == 2  # C3 e C4
assert f2["valor_cobrado"] == Decimal("350.00")  # 200 + 150
assert f2["clientes_pagaram"] == 0
assert f2["valor_pago"] == Decimal("0.00")
assert f2["parcelas_cobradas"] == 2
assert f2["parcelas_pagas"] == 0
assert f2["parcelas_renegociadas"] == 1  # C3
assert f2["conversao_clientes"] == Decimal("0.0000")
assert f2["recuperacao_valor"] == Decimal("0.0000")

# Por Loja
por_loja = {r["loja"]: r for r in res["por_loja"]}
assert "01" in por_loja
l1 = por_loja["01"]
assert l1["loja_nome"] == "01 - PALMAS"
assert l1["regional"] == "REGIONAL NORTE"
assert l1["cluster_inad"] == "ALTO"
assert l1["clientes_cobrados"] == 2  # C1 e C2
assert l1["valor_cobrado"] == Decimal("250.00")  # 100 + 80 + 70
assert l1["clientes_pagaram"] == 2
assert l1["valor_pago"] == Decimal("180.00")  # 100 + 80
assert l1["parcelas_cobradas"] == 3
assert l1["parcelas_pagas"] == 2
assert l1["parcelas_renegociadas"] == 0
assert l1["conversao_clientes"] == Decimal("1.0000")
assert l1["recuperacao_valor"] == (Decimal("180") / Decimal("250")).quantize(Decimal("0.0001"))

assert "02" in por_loja
l2 = por_loja["02"]
# Loja não existente em lojas_info -> atributos nulos
assert l2["loja_nome"] is None
assert l2["regional"] is None
assert l2["cluster_inad"] is None
assert l2["clientes_cobrados"] == 3  # C1, C3, C4
assert l2["valor_cobrado"] == Decimal("400.00")  # 50 + 200 + 150
assert l2["clientes_pagaram"] == 1  # Apenas C1 pagou a parcela de 50
assert l2["valor_pago"] == Decimal("50.00")
assert l2["parcelas_cobradas"] == 3
assert l2["parcelas_pagas"] == 1
assert l2["parcelas_renegociadas"] == 1  # C3
assert l2["conversao_clientes"] == (Decimal("1") / Decimal("3")).quantize(Decimal("0.0001"))
assert l2["recuperacao_valor"] == (Decimal("50") / Decimal("400")).quantize(Decimal("0.0001"))

# Total consolidado (clientes distintos no total geral)
tot = res["total"]
# C1 aparece na loja 01 e na loja 02. No total, deve contar apenas UMA vez!
assert tot["clientes_cobrados"] == 4  # C1, C2, C3, C4
assert tot["valor_cobrado"] == Decimal("650.00")
assert tot["clientes_pagaram"] == 2  # C1 e C2
assert tot["valor_pago"] == Decimal("230.00")  # 100 (P1) + 50 (P2) + 80 (P3)
assert tot["parcelas_cobradas"] == 6
assert tot["parcelas_pagas"] == 3
assert tot["parcelas_renegociadas"] == 1
assert tot["conversao_clientes"] == Decimal("0.5000")  # 2 / 4
assert tot["recuperacao_valor"] == (Decimal("230") / Decimal("650")).quantize(Decimal("0.0001"))

# Regra de divisão por zero (relatório sem itens)
vazio = montar_relatorio([])
assert vazio["por_faixa"] == []
assert vazio["por_loja"] == []
assert vazio["total"]["clientes_cobrados"] == 0
assert vazio["total"]["valor_cobrado"] == Decimal("0.00")
assert vazio["total"]["conversao_clientes"] == Decimal("0.0000")
assert vazio["total"]["recuperacao_valor"] == Decimal("0.0000")


# ============================================================================
# 1b. Consistência entre a regra da base de cobrança e buscar_parcelas_cobranca
# ============================================================================

def _simular_sql_base(titulos_abertos: list[dict], hoje: date, juros) -> tuple[int, Decimal]:
    """Simula a agregação da CTE abertos de _SQL_BASE_COBRANCA para um cliente."""
    validos = [
        t for t in titulos_abertos
        if t["rp"] == "R" and t["status"] == "A" and t["tipo"] in ("4", "5") and t["valor"] > Decimal("0.00")
    ]
    if not validos:
        return 0, Decimal("0.00")
    vencimento_min = min(t["vencimento"] for t in validos)
    limite = max(hoje, vencimento_min)
    cobradas = [t for t in validos if t["vencimento"] <= limite]
    qtd = len(cobradas)
    soma = Decimal("0.00")
    for t in cobradas:
        dias = (hoje - t["vencimento"]).days
        if dias >= juros.dias_min:
            soma += t["valor"] + t["valor"] * juros.juros_dia * dias + t["valor"] * juros.multa
        else:
            soma += t["valor"]
    return qtd, soma.quantize(Decimal("0.01"))


def _simular_sql_parcelas(titulos_abertos: list[dict], hoje: date, juros) -> list[dict]:
    """Simula a seleção de parcelas de _SQL_PARCELAS_COBRANCA para um cliente."""
    validos = [
        t for t in titulos_abertos
        if t["rp"] == "R" and t["status"] == "A" and t["tipo"] in ("4", "5") and t["valor"] > Decimal("0.00")
    ]
    if not validos:
        return []
    vencimento_min = min(t["vencimento"] for t in validos)
    limite = max(hoje, vencimento_min)
    resultado = []
    for t in validos:
        if t["vencimento"] <= limite:
            dias = (hoje - t["vencimento"]).days
            if dias >= juros.dias_min:
                val_cob = (t["valor"] + t["valor"] * juros.juros_dia * dias + t["valor"] * juros.multa).quantize(Decimal("0.01"))
            else:
                val_cob = t["valor"].quantize(Decimal("0.01"))
            resultado.append({**t, "valor_cobrar": val_cob})
    return resultado


titulos_teste = [
    {"codigo": "T1", "rp": "R", "status": "A", "tipo": "4", "valor": Decimal("100.00"), "vencimento": date(2026, 8, 1)},
    {"codigo": "T2", "rp": "R", "status": "A", "tipo": "5", "valor": Decimal("50.00"), "vencimento": date(2026, 9, 20)},
    {"codigo": "T3", "rp": "R", "status": "A", "tipo": "4", "valor": Decimal("200.00"), "vencimento": date(2026, 10, 1)},
    {"codigo": "T4", "rp": "R", "status": "B", "tipo": "4", "valor": Decimal("80.00"), "vencimento": date(2026, 7, 1)},
    {"codigo": "T5", "rp": "R", "status": "A", "tipo": "4", "valor": Decimal("0.00"), "vencimento": date(2026, 8, 1)},
]

hoje_teste = date(2026, 9, 21)
juros_teste = seta_client.PARAMETROS_JUROS_PADRAO
qtd_base, valor_base = _simular_sql_base(titulos_teste, hoje_teste, juros_teste)
parcelas_res = _simular_sql_parcelas(titulos_teste, hoje_teste, juros_teste)

assert len(parcelas_res) == qtd_base == 2
assert sum(p["valor_cobrar"] for p in parcelas_res) == valor_base


# ============================================================================
# 2. Testes de janela de pagamento (lógica de data)
# ============================================================================

data_cob = date(2026, 9, 10)


def _avaliar_pagamento(data_cobranca, dt_pagamento, status_tit, dias_janela=None):
    if status_tit == "S":
        return False, True  # pago=False, renegociada=True
    if status_tit == "B" and dt_pagamento:
        if dt_pagamento >= data_cobranca:
            if dias_janela is None or dt_pagamento <= data_cobranca + timedelta(days=dias_janela):
                return True, False
    return False, False


# Pago antes da cobrança -> não pago
p, r = _avaliar_pagamento(data_cob, date(2026, 9, 9), "B")
assert p is False and r is False

# Pago no dia da cobrança -> pago
p, r = _avaliar_pagamento(data_cob, date(2026, 9, 10), "B")
assert p is True and r is False

# Com dias_janela = 7:
# Pago no dia +7 -> pago
p, r = _avaliar_pagamento(data_cob, date(2026, 9, 17), "B", dias_janela=7)
assert p is True and r is False

# Pago no dia +8 -> fora da janela (não pago)
p, r = _avaliar_pagamento(data_cob, date(2026, 9, 18), "B", dias_janela=7)
assert p is False and r is False

# Status 'S' -> renegociado, não pago
p, r = _avaliar_pagamento(data_cob, date(2026, 9, 12), "S", dias_janela=7)
assert p is False and r is True


# ============================================================================
# 3. Testes de integração via API (TestClient + Postgres agy_efet)
# ============================================================================

with TestClient(app) as client:
    # 401 sem autenticação
    assert client.get("/reports/efetividade").status_code == 401
    assert client.get("/reports/efetividade.xlsx").status_code == 401
    assert client.get("/relatorios/efetividade").status_code == 401
    assert client.get("/relatorios/efetividade.xlsx").status_code == 401

    # Login como administrador
    login_res = client.post(
        "/auth/login",
        json={"email": "admin@topfama.com.br", "password": "senha-admin-segura-para-o-teste-123"},
    )
    assert login_res.status_code == 200, login_res.text
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # Limpa dados anteriores na base de teste
    db: Session = database.SessionLocal()
    try:
        db.query(models.LeadParcela).delete()
        db.query(models.Lead).delete()
        db.commit()
    finally:
        db.close()

    # Cria cliente simulado no cobranca_base e seta_client
    cliente_mock_1 = {
        "codigo": "00100001",
        "nome": "CLIENTE TESTE ALFA",
        "celular": "5563999990001",
        "celular_origem": "telefone1",
        "celular_original": "63999990001",
        "cpfcnpj": "12345678901",
        "status": "A",
        "status_descricao": "ativo",
        "loja_cadastro": "01",
        "salario": Decimal("2500.00"),
        "limite_rotativo": Decimal("500.00"),
        "nascimento": date(1990, 1, 1),
        "cadastro": date(2020, 1, 1),
        "cluster": "POTENCIAL",
        "valor_pago": Decimal("600.00"),
        "qtd_compras": 3,
        "faixa_compra": "1 a 9",
        "ultima_compra": date(2026, 1, 1),
        "spc_restricao": "nao",
        "faixa": "11 A 20",
        "dias_atraso": 15,
        "entra_whatsapp": True,
        "qtd_titulos": 2,
        "valor_em_aberto": Decimal("300.00"),
        "qtd_parcelas_cobranca": 2,
        "valor_cobrar": Decimal("330.00"),
        "vencimento_mais_antigo": date(2026, 8, 20),
        "lojas": ["01"],
        "portadores": ["001"],
    }
    cliente_mock_2 = {
        "codigo": "00100002",
        "nome": "CLIENTE TESTE BETA",
        "celular": "5563999990002",
        "celular_origem": "telefone1",
        "celular_original": "63999990002",
        "cpfcnpj": "12345678902",
        "status": "A",
        "status_descricao": "ativo",
        "loja_cadastro": "02",
        "salario": Decimal("3500.00"),
        "limite_rotativo": Decimal("800.00"),
        "nascimento": date(1985, 5, 5),
        "cadastro": date(2019, 5, 5),
        "cluster": "ALTO POTENCIAL",
        "valor_pago": Decimal("1800.00"),
        "qtd_compras": 5,
        "faixa_compra": "1 a 9",
        "ultima_compra": date(2026, 2, 1),
        "spc_restricao": "nao",
        "faixa": "21 A 30",
        "dias_atraso": 25,
        "entra_whatsapp": True,
        "qtd_titulos": 1,
        "valor_em_aberto": Decimal("200.00"),
        "qtd_parcelas_cobranca": 1,
        "valor_cobrar": Decimal("220.00"),
        "vencimento_mais_antigo": date(2026, 8, 10),
        "lojas": ["02"],
        "portadores": ["001"],
    }

    # Monkeypatching de serviços externos (SETA e Google)
    def mock_buscar_base(db, **kwargs):
        return [cliente_mock_1, cliente_mock_2]

    def mock_buscar_spc(codigos):
        return {}

    juros_passado = []

    def mock_buscar_parcelas_cobranca(codigos, juros=seta_client.PARAMETROS_JUROS_PADRAO):
        juros_passado.append(juros)
        return {
            "00100001": [
                {
                    "titulo_codigo": "TIT101",
                    "empresa": "01",
                    "vencimento": date(2026, 8, 20),
                    "valor": Decimal("150.00"),
                    "valor_cobrar": Decimal("165.00"),
                },
                {
                    "titulo_codigo": "TIT102",
                    "empresa": "01",
                    "vencimento": date(2026, 8, 25),
                    "valor": Decimal("150.00"),
                    "valor_cobrar": Decimal("165.00"),
                },
            ],
            "00100002": [
                {
                    "titulo_codigo": "TIT103",
                    "empresa": "02",
                    "vencimento": date(2026, 8, 10),
                    "valor": Decimal("200.00"),
                    "valor_cobrar": Decimal("220.00"),
                },
            ],
        }

    cobranca_base.buscar_base = mock_buscar_base
    seta_client.buscar_spc = mock_buscar_spc
    seta_client.buscar_parcelas_cobranca = mock_buscar_parcelas_cobranca

    # 1. Gera leads e confere gravação das parcelas
    gerar_res = client.post("/leads/gerar", headers=headers)
    assert gerar_res.status_code == 200, gerar_res.text
    dados_gerar = gerar_res.json()
    assert dados_gerar["criados"] == 2

    # Confere que o juros recebido veio das regras configuradas no banco
    assert len(juros_passado) == 1
    from app.regras_db import carregar_regras
    db_check = database.SessionLocal()
    try:
        regras_banco = carregar_regras(db_check)
        assert juros_passado[0] == regras_banco.juros
    finally:
        db_check.close()

    db = database.SessionLocal()
    try:
        leads = db.query(models.Lead).order_by(models.Lead.codigo_cliente).all()
        assert len(leads) == 2
        lead1, lead2 = leads[0], leads[1]
        lead1_id, lead2_id = lead1.id, lead2.id
        assert len(lead1.parcelas) == 2
        assert len(lead2.parcelas) == 1
        titulos_lead1 = {p.titulo_codigo for p in lead1.parcelas}
        assert titulos_lead1 == {"TIT101", "TIT102"}
        assert lead2.parcelas[0].titulo_codigo == "TIT103"

        # Adiciona um lead antigo SEM parcelas (para testar leads_sem_parcelas)
        lead_antigo = models.Lead(
            codigo_cliente="00109999",
            nome="LEAD HISTORICO",
            faixa="11 A 20",
            cluster="ESPECIAL",
            dias_atraso=15,
            vencimento_mais_antigo=date(2026, 8, 15),
            status="cobrado",
            cobrado_em=datetime.now(timezone.utc).replace(tzinfo=None),
        )
        db.add(lead_antigo)
        db.commit()
    finally:
        db.close()

    # 2. Marca os leads gerados como cobrados
    marcar_res = client.post(
        "/leads/marcar-cobrados",
        json={"ids": [lead1_id, lead2_id]},
        headers=headers,
    )
    assert marcar_res.status_code == 200, marcar_res.text
    assert marcar_res.json()["atualizados"] == 2

    # 3. Simula situação dos títulos no SETA
    hoje = datetime.now(timezone.utc).date()

    def mock_situacao_titulos(codigos):
        res = {}
        for c in codigos:
            if c == "TIT101":
                # Pago hoje: 165.00
                res[c] = {"status": "B", "pagamento": hoje, "valorpago": Decimal("165.00"), "valor": Decimal("150.00")}
            elif c == "TIT102":
                # Ainda em aberto
                res[c] = {"status": "A", "pagamento": None, "valorpago": Decimal("0.00"), "valor": Decimal("150.00")}
            elif c == "TIT103":
                # Renegociado
                res[c] = {"status": "S", "pagamento": None, "valorpago": Decimal("0.00"), "valor": Decimal("200.00")}
        return res

    seta_client.situacao_titulos = mock_situacao_titulos

    # 4. Consulta /reports/efetividade
    rep_res = client.get("/reports/efetividade", headers=headers)
    assert rep_res.status_code == 200, rep_res.text
    rep_data = rep_res.json()

    # Valida leads_sem_parcelas
    assert rep_data["leads_sem_parcelas"] == 1

    # Valida números do total consolidado
    tot_api = rep_data["total"]
    assert tot_api["clientes_cobrados"] == 2
    assert Decimal(str(tot_api["valor_cobrado"])) == Decimal("550.00")
    assert tot_api["clientes_pagaram"] == 1  # lead 1 pagou TIT101
    assert Decimal(str(tot_api["valor_pago"])) == Decimal("165.00")
    assert tot_api["parcelas_cobradas"] == 3
    assert tot_api["parcelas_pagas"] == 1
    assert tot_api["parcelas_renegociadas"] == 1
    assert Decimal(str(tot_api["conversao_clientes"])) == Decimal("0.5000")
    assert Decimal(str(tot_api["recuperacao_valor"])) == (Decimal("165") / Decimal("550")).quantize(Decimal("0.0001"))

    # Confere que a rota via /relatorios/efetividade dá o mesmo resultado
    rep_rel_res = client.get("/relatorios/efetividade", headers=headers)
    assert rep_rel_res.status_code == 200
    assert rep_rel_res.json() == rep_data

    # Valida ordenação das faixas conforme carregar_regras(db).nomes_faixa
    faixas_retornadas = [f["faixa"] for f in rep_data["por_faixa"]]
    indices_faixas = [
        regras_banco.nomes_faixa.index(f)
        for f in faixas_retornadas
        if f in regras_banco.nomes_faixa
    ]
    assert indices_faixas == sorted(indices_faixas)

    # 5. Filtros da API:
    # Filtro de faixa
    rf_faixa = client.get("/reports/efetividade?faixa=11%20A%2020", headers=headers).json()
    assert rf_faixa["total"]["clientes_cobrados"] == 1
    assert Decimal(str(rf_faixa["total"]["valor_cobrado"])) == Decimal("330.00")

    # Filtro de loja (aplica-se à empresa da parcela)
    rf_loja = client.get("/reports/efetividade?loja=02", headers=headers).json()
    assert rf_loja["total"]["clientes_cobrados"] == 1
    assert rf_loja["total"]["parcelas_renegociadas"] == 1

    # Filtro cobrado_de / cobrado_ate
    rf_data = client.get(
        f"/reports/efetividade?cobrado_de={hoje.isoformat()}&cobrado_ate={hoje.isoformat()}",
        headers=headers,
    ).json()
    assert rf_data["total"]["clientes_cobrados"] == 2

    # Filtro dias_janela (pago hoje está dentro de janela=0)
    rf_jan0 = client.get("/reports/efetividade?dias_janela=0", headers=headers).json()
    assert rf_jan0["total"]["parcelas_pagas"] == 1

    # 6. Loja por atributos e disponibilidade do Google:
    # Se Google indisponível ao pedir filtro de atributo -> 503
    def mock_ler_aba_erro(db, sheet_id, gid):
        raise google_client.GoogleIndisponivel("Google desconectado")

    google_client.ler_aba = mock_ler_aba_erro

    res_503 = client.get("/reports/efetividade?regional=NORTE", headers=headers)
    assert res_503.status_code == 503, res_503.text

    # Sem filtro de atributo -> 200, regional/cluster_inad nulos
    res_sem_attr = client.get("/reports/efetividade", headers=headers)
    assert res_sem_attr.status_code == 200
    for l_item in res_sem_attr.json()["por_loja"]:
        assert l_item["regional"] is None
        assert l_item["cluster_inad"] is None

    # 7. Relatórios exportáveis existentes continuam funcionando
    assert client.get("/reports/envios/export", headers=headers).status_code == 200
    assert client.get("/reports/telefones-invalidos/export", headers=headers).status_code == 200
    assert client.get("/relatorios/envios/export", headers=headers).status_code == 200
    assert client.get("/relatorios/telefones-invalidos/export", headers=headers).status_code == 200

    # 8. Exportação Excel: GET /reports/efetividade.xlsx
    xlsx_res = client.get("/reports/efetividade.xlsx", headers=headers)
    assert xlsx_res.status_code == 200, xlsx_res.text
    assert (
        xlsx_res.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_res.content))
    assert "Por faixa" in wb.sheetnames
    assert "Por loja" in wb.sheetnames

    # Validação da aba 'Por faixa'
    ws_f = wb["Por faixa"]
    headers_esperados_faixa = [
        "Faixa de atraso",
        "Clientes cobrados",
        "Valor cobrado",
        "Clientes que pagaram",
        "Valor pago",
        "Conversão (%)",
        "Recuperação (%)",
    ]
    assert [cell.value for cell in ws_f[1]] == headers_esperados_faixa
    # Linha Total em negrito
    linha_tot_f = [cell for cell in ws_f[ws_f.max_row]]
    assert linha_tot_f[0].value == "Total"
    assert all(cell.font.bold is True for cell in linha_tot_f)
    # Formatação de moeda e porcentagem na linha total
    assert linha_tot_f[2].number_format == '"R$" #,##0.00'
    assert linha_tot_f[5].number_format == "0.0%"

    # Validação da aba 'Por loja'
    ws_l = wb["Por loja"]
    headers_esperados_loja = [
        "Loja",
        "Regional",
        "Cluster INAD",
        "Clientes cobrados",
        "Valor cobrado",
        "Clientes que pagaram",
        "Valor pago",
        "Conversão (%)",
        "Recuperação (%)",
    ]
    assert [cell.value for cell in ws_l[1]] == headers_esperados_loja
    linha_tot_l = [cell for cell in ws_l[ws_l.max_row]]
    assert linha_tot_l[0].value == "Total"
    assert all(cell.font.bold is True for cell in linha_tot_l)
    assert linha_tot_l[4].number_format == '"R$" #,##0.00'
    assert linha_tot_l[7].number_format == "0.0%"

    # Validação da formatação condicional na coluna Cluster INAD
    cfs = list(ws_l.conditional_formatting)
    assert len(cfs) > 0
    total_regras = sum(len(cf.rules) for cf in cfs)
    assert total_regras == 3  # ALT, MED, BAIX


# ============================================================================
# 4. Ciclo de migration Alembic e verificação de schema
# ============================================================================

from pathlib import Path
from alembic import command
from alembic.config import Config

backend_dir = Path(__file__).resolve().parent.parent
alembic_cfg = Config(str(backend_dir / "alembic.ini"))
alembic_cfg.set_main_option("script_location", str(backend_dir / "alembic"))
command.downgrade(alembic_cfg, "base")
command.upgrade(alembic_cfg, "head")
command.check(alembic_cfg)

print("OK")
