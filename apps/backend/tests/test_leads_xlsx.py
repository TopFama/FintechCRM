"""Script de validação do módulo leads_xlsx — executa com assert simples,
sem pytest. Encerra imprimindo 'OK'."""

import io
import time
import types

import openpyxl

from app.utils.leads_xlsx import (
    formatar_codigo,
    formatar_celular,
    formatar_cpf,
    gerar_xlsx_leads,
    linhas_para_exportacao,
    primeiro_nome,
)


# formatar_codigo

assert formatar_codigo("123456") == "00123456"
assert formatar_codigo("00123456") == "00123456"
assert formatar_codigo(" 1234 ") == "00001234"
assert formatar_codigo("12.345") == "00012345"
assert formatar_codigo("123456789") == "123456789"
assert formatar_codigo("") == ""
assert formatar_codigo(None) == ""


# primeiro_nome

assert primeiro_nome("MARIA DA SILVA SANTOS") == "Maria"
assert primeiro_nome("  joão   pedro ") == "João"
assert primeiro_nome("JOSÉ") == "José"
assert primeiro_nome("ANA-LUIZA PEREIRA") == "Ana-Luiza"
assert primeiro_nome("D'AVILA SOUZA") == "D'Avila"
assert primeiro_nome("MARIA2 SILVA") == "Maria"
assert primeiro_nome("123 456") == ""
assert primeiro_nome("") == ""
assert primeiro_nome(None) == ""
# Novos casos (REVISAO.md)
assert primeiro_nome("MARIA- SILVA") == "Maria", repr(primeiro_nome("MARIA- SILVA"))
assert primeiro_nome("-JOAO") == "Joao", repr(primeiro_nome("-JOAO"))
assert primeiro_nome("'ANA") == "Ana", repr(primeiro_nome("'ANA"))
assert primeiro_nome("MARIA_JOSE") == "Maria", repr(primeiro_nome("MARIA_JOSE"))
assert primeiro_nome("J. SILVA") == "J", repr(primeiro_nome("J. SILVA"))


# formatar_cpf

assert formatar_cpf("52998224725") == "529.982.247-25"
assert formatar_cpf("529.982.247-25") == "529.982.247-25"
assert formatar_cpf("5998224725") == "059.982.247-25"
assert formatar_cpf("998224725") == "009.982.247-25"
assert formatar_cpf("11222333000181") == "11.222.333/0001-81"
assert formatar_cpf("1222333000181") == "01.222.333/0001-81"
assert formatar_cpf("123456789012345") == "123456789012345"
assert formatar_cpf("") == ""
assert formatar_cpf(None) == ""


# formatar_celular

assert formatar_celular("5563991234567") == "5563991234567"
assert formatar_celular("(63) 99123-4567") == "5563991234567"
assert formatar_celular("63991234567") == "5563991234567"
assert formatar_celular("5500000000000") == ""
assert formatar_celular("123") == ""
assert formatar_celular("") == ""
assert formatar_celular(None) == ""


# gerar_xlsx_leads

def _lead(codigo, nome, cpf, celular):
    return types.SimpleNamespace(
        codigo_cliente=codigo,
        nome=nome,
        cpf=cpf,
        celular=celular,
    )


# Três leads: cpf sem máscara/zeros, celular None, nome completo em maiúsculas
leads = [
    _lead("00123456", "MARIA DA SILVA SANTOS", "52998224725", "63991234567"),
    _lead("456", "PEDRO ALVES", "998224725", None),
    _lead("789", "ANA-LUIZA PEREIRA", "11222333000181", "(63) 99123-4567"),
]

bytes_ = gerar_xlsx_leads(leads)
assert isinstance(bytes_, bytes)
assert len(bytes_) > 0

wb = openpyxl.load_workbook(io.BytesIO(bytes_))

# Nome da aba
assert "Leads" in wb.sheetnames
ws = wb["Leads"]

# Cabeçalho exato
assert [ws.cell(1, c).value for c in range(1, 5)] == ["Codigo", "Nome", "Celular", "CPF"]

# Cor de preenchimento do cabeçalho termina em FCE4D6
for col in range(1, 5):
    rgb = ws.cell(1, col).fill.fgColor.rgb
    assert rgb.endswith("FCE4D6"), f"cor incorreta na coluna {col}: {rgb}"

# Fonte do cabeçalho não é negrito
for col in range(1, 5):
    assert not ws.cell(1, col).font.bold, f"coluna {col} com negrito inesperado"

# Painel congelado
assert ws.freeze_panes == "A2"

# Larguras de coluna
from openpyxl.utils import get_column_letter
larguras_esperadas = {"A": 12, "B": 24, "C": 18, "D": 16}
for letra, esperada in larguras_esperadas.items():
    real = ws.column_dimensions[letra].width
    assert real == esperada, f"largura coluna {letra}: esperado {esperada}, obteve {real}"

# Linhas de dados — valores formatados e number_format "@"
linha2 = [ws.cell(2, c) for c in range(1, 5)]
assert linha2[0].value == "00123456"
assert linha2[1].value == "Maria"
assert linha2[2].value == "5563991234567"
assert linha2[3].value == "529.982.247-25"
for cell in linha2:
    assert cell.number_format == "@", f"number_format incorreto: {cell.number_format}"

linha3 = [ws.cell(3, c) for c in range(1, 5)]
assert linha3[0].value == "00000456"
assert linha3[1].value == "Pedro"
assert linha3[3].value == "009.982.247-25"
assert linha3[2].value in ("", None)   # celular None → célula vazia (openpyxl relê "" como None)
for cell in linha3:
    assert cell.number_format == "@"

linha4 = [ws.cell(4, c) for c in range(1, 5)]
assert linha4[0].value == "00000789"
assert linha4[1].value == "Ana-Luiza"
assert linha4[2].value == "5563991234567"
assert linha4[3].value == "11.222.333/0001-81"
for cell in linha4:
    assert cell.number_format == "@"

# Ordem de entrada preservada
linhas = linhas_para_exportacao(leads)
assert linhas[0][1] == "Maria"
assert linhas[1][1] == "Pedro"
assert linhas[2][1] == "Ana-Luiza"

# Lista vazia → só cabeçalho (1 linha)
bytes_vazio = gerar_xlsx_leads([])
ws_vazio = openpyxl.load_workbook(io.BytesIO(bytes_vazio))["Leads"]
assert ws_vazio.max_row == 1
assert [ws_vazio.cell(1, c).value for c in range(1, 5)] == ["Codigo", "Nome", "Celular", "CPF"]


# Desempenho: 20.000 leads em menos de 10 s e planilha com 20.001 linhas
leads_grandes = [
    _lead(f"{i:08d}", f"NOME SOBRENOME {i}", "52998224725", "63991234567")
    for i in range(20_000)
]
t0 = time.perf_counter()
bytes_grandes = gerar_xlsx_leads(leads_grandes)
elapsed = time.perf_counter() - t0
print(f"20.000 leads gerados em {elapsed:.2f} s")
assert elapsed < 10, f"geração lenta: {elapsed:.2f} s"
ws_grandes = openpyxl.load_workbook(io.BytesIO(bytes_grandes))["Leads"]
assert ws_grandes.max_row == 20_001, f"linhas: {ws_grandes.max_row}"

print("OK")
