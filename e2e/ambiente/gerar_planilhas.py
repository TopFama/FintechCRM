"""Gera as planilhas fictícias usadas no teste de importação (tests/dados/).
Rode de novo só se mudar o conteúdo: python ambiente/gerar_planilhas.py"""

from pathlib import Path

from openpyxl import Workbook

DESTINO = Path(__file__).resolve().parents[1] / "tests" / "dados"
DESTINO.mkdir(parents=True, exist_ok=True)

wb = Workbook()
ws = wb.active
ws.append(["Codigo Cliente", "Nome Completo", "Documento", "Telefone", "Valor Devido", "Vencimento"])
ws.append(["101", "Roberta Almeida Prado", "12345678909", "(11) 98888-1111", "150,00", "10/09/2026"])
ws.append(["102", "Sérgio Tavares", "98765432100", "11977772222", "89,90", "12/09/2026"])
ws.append(["103", "Helena Duarte", "11144477735", "(21) 96666-3333", "300,50", "15/09/2026"])
ws.append(["104", "Otávio Brandão", "22233344405", "123", "45,00", "15/09/2026"])        # telefone inválido
ws.append(["105", "Márcia Quintela", "33322211100", "(11) 3333-4444", "60,00", "16/09/2026"])  # fixo: inválido
ws.append(["106", "Lúcia Ferraz", "55566677788", "(11) 94444-7777", "", ""])                  # variável em branco: entra na fila como erro
ws.append(["", "Sem Código", "44455566600", "(11) 95555-6666", "10,00", "16/09/2026"])          # sem código
ws.append(["107", "Paulo Zanetti", "66677788899", "(11) 97777-0107", "0,00", "16/09/2026"])    # valor zerado: entra na fila como erro
wb.save(DESTINO / "planilha_faixa.xlsx")

# Campanhas: relatório de lojas (coluna de código) e planilha de clientes
wb = Workbook()
ws = wb.active
ws.append(["FILIAL", "Observação"])
ws.append([1, "centro"])
ws.append([6, "praia"])
ws.append([99, "não existe"])
wb.save(DESTINO / "lojas_campanha.xlsx")

wb = Workbook()
ws = wb.active
ws.append(["Codigo", "Nome", "Valor", "Obs"])
ws.append([27, "Cliente A", "99,90", "Promo A"])
ws.append([36, "Cliente B", "150,00", "Promo B"])
ws.append([99999, "Não existe no SETA", "10,00", "Promo C"])
ws.append(["", "Sem código", "10,00", "x"])
wb.save(DESTINO / "clientes_campanha.xlsx")

(DESTINO / "nao_e_planilha.xlsx").write_text("isto não é um arquivo Excel\n")
print("planilhas geradas em", DESTINO)
